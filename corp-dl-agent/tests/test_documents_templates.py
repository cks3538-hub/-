"""documents.templates: fill_pptx / fill_xlsx — 승인 placeholder·named range 만 수정, 구조·수식 보존, 호환성 보고."""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import pytest

from corp_dl_agent.common import Quantity, sha256_file
from corp_dl_agent.documents.payload import PayloadItem
from corp_dl_agent.documents.synthetic import example_payload
from corp_dl_agent.documents.templates import (
    collect_pptx_texts,
    collect_xlsx_texts,
    fill_pptx,
    fill_xlsx,
    read_pptx_structure,
    read_xlsx_structure,
)
from corp_dl_agent.errors import AgentError

pytestmark = pytest.mark.documents

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "documents"
PPTX_TPL = FIXTURES / "template_review.pptx"
XLSX_TPL = FIXTURES / "template_comparison.xlsx"


def test_fill_pptx_replaces_only_placeholders_and_preserves_structure(tmp_path: Path) -> None:
    payload = example_payload()
    before = sha256_file(PPTX_TPL)
    out = tmp_path / "산출" / "review.pptx"
    m = fill_pptx(PPTX_TPL, payload, out)
    assert out.is_file() and sha256_file(PPTX_TPL) == before and m.template_hash_unchanged
    assert m.output_hash == sha256_file(out) and m.input_hash == before
    assert m.recalc_status == "RECALC_NOT_RUN" and m.render_status == "RENDER_NOT_RUN" and m.synthetic
    assert m.missing_keys == [] and m.unknown_keys == [] and m.unsupported_elements == []
    texts = collect_pptx_texts(out)
    assert texts["slide:1/shape:Title 1"] == payload.subject
    body = texts["slide:2/shape:Body 1"]
    assert "중량 A: 11.2 kg" in body and "중량 차이(B-A): -0.8 kg" in body and "차종: X-NEW (synthetic)" in body
    # run 이 나뉜 placeholder 도 치환되고 서식(run) 은 유지된다
    assert "원가 합계: 1,980 원 (출처 cost:asm_b:total)" in body and "{{" not in body
    assert texts["slide:3/table:Table 1/r2c2"] == "11.2 kg" and texts["slide:3/table:Table 1/r3c4"] == "-70 원"
    assert texts["slide:3/table:Table 1/r2c1"] == "중량 A"
    assert texts["slide:3/notes"] == "근거 요약: evidence.json 참조 (synthetic 템플릿)"
    assert "{{" not in "".join(texts.values())
    tpl = read_pptx_structure(PPTX_TPL)
    cur = read_pptx_structure(out)
    assert tpl == cur
    by_key = {(r.locator, r.key, r.attr): r for r in m.replacements}
    assert by_key[("slide:2/shape:Body 1", "cost_total", None)].rendered == "1,980 원"
    assert by_key[("slide:2/shape:Body 1", "cost_total", None)].numeric and by_key[
        ("slide:2/shape:Body 1", "cost_total", None)
    ].value == 1980
    assert by_key[("slide:3/table:Table 1/r2c2", "mass_a", "unit")].rendered == "kg"
    from pptx import Presentation

    prs = Presentation(str(out))
    runs = prs.slides[1].shapes[1].text_frame.paragraphs[5].runs
    assert runs[0].font.bold is True and runs[0].text.startswith("원가 합계: 1,980 원")


def test_fill_pptx_missing_conflict_unknown_and_unapproved(tmp_path: Path) -> None:
    payload = example_payload()
    payload.items["mass_b"].quantity = Quantity.missing("kg")
    payload.items["cost_b"].quantity = Quantity.conflict("원", "견적 vs 계산 상충")
    del payload.texts["conclusion"]
    m = fill_pptx(PPTX_TPL, payload, tmp_path / "review.pptx")
    texts = collect_pptx_texts(tmp_path / "review.pptx")
    assert "중량 B: MISSING" in texts["slide:2/shape:Body 1"] and "결론: MISSING" in texts["slide:2/shape:Body 1"]
    assert texts["slide:3/table:Table 1/r3c3"] == "CONFLICT"
    assert m.missing_keys == ["cost_b", "mass_b"] and m.unknown_keys == ["conclusion"]
    m2 = fill_pptx(PPTX_TPL, example_payload(), tmp_path / "review2.pptx", approved_keys={"subject", "mass_a"})
    texts2 = collect_pptx_texts(tmp_path / "review2.pptx")
    assert "{{mass_b}}" in texts2["slide:2/shape:Body 1"] and "중량 A: 11.2 kg" in texts2["slide:2/shape:Body 1"]
    assert "mass_b" in m2.unapproved_keys and "subject" not in m2.unapproved_keys


def test_fill_pptx_preserves_chart_and_reports_unsupported(tmp_path: Path) -> None:
    out = tmp_path / "past_copy.pptx"
    m = fill_pptx(FIXTURES / "past_review_2023.pptx", example_payload(), out)
    assert [u.element_type for u in m.unsupported_elements] == ["chart"]
    assert m.unsupported_elements[0].preserved and m.unsupported_elements[0].locator == "slide:3/chart:Chart 1"
    assert m.changed_locators == [] and read_pptx_structure(out)["n_charts"] == 1
    assert read_pptx_structure(out) == read_pptx_structure(FIXTURES / "past_review_2023.pptx")
    with pytest.raises(AgentError) as ei:
        fill_pptx(PPTX_TPL, example_payload(), PPTX_TPL)
    assert ei.value.code == "E_INPUT_INVALID"


def test_fill_xlsx_named_ranges_formulas_preserved_recalc_not_run(tmp_path: Path) -> None:
    import openpyxl

    payload = example_payload()
    before = sha256_file(XLSX_TPL)
    out = tmp_path / "산출" / "comparison.xlsx"
    m = fill_xlsx(XLSX_TPL, payload, out)
    assert sha256_file(XLSX_TPL) == before and m.template_hash_unchanged and m.recalc_status == "RECALC_NOT_RUN"
    tpl = read_xlsx_structure(XLSX_TPL)
    cur = read_xlsx_structure(out)
    assert tpl == cur
    # 템플릿의 수식 셀 6개가 원문 그대로 보존된다 (차이 3 + 합계 2 + Lines 합계 1)
    assert sorted(cur["formulas"]) == [
        "Comparison!B16",
        "Comparison!B18",
        "Comparison!D10",
        "Comparison!D8",
        "Comparison!D9",
        "Lines!B8",
    ]
    assert cur["formulas"]["Comparison!D8"] == "=C8-B8" and cur["formulas"]["Comparison!B16"] == "=SUM(B13:B15)"
    wb = openpyxl.load_workbook(out)
    ws = wb["Comparison"]
    assert ws["B8"].value == 11.2 and ws["C8"].value == 10.4 and ws["E8"].value == "kg"
    assert ws["B9"].value == 2050 and ws["B17"].value == 1980 and ws["B4"].value == "X-NEW (synthetic)"
    assert ws["B1"].value == payload.subject and ws["B20"].value == "synthetic"
    assert ws["B16"].value == "=SUM(B13:B15)" and ws["B8"].number_format == "#,##0.##"
    lines = wb["Lines"]
    assert lines["A2"].value == "재료비" and lines["B2"].value == 900 and lines["B4"].value == 300
    assert lines["A5"].value is None and lines["B8"].value == "=SUM(B2:B6)"
    # cached value 없음 → RECALC_NOT_RUN
    cached = openpyxl.load_workbook(out, data_only=True)["Comparison"]["B16"].value
    assert cached is None
    assert "range:cost_total_formula" not in m.changed_locators
    assert {u.element_type for u in m.unsupported_elements} == {"formula_cell"}
    assert "range:tbl_cost_lines/r1c2" in m.changed_locators and m.missing_keys == [] and m.unknown_keys == []
    texts = collect_xlsx_texts(out)
    assert texts["range:mass_a"] == "11.2" and texts["sheet:Comparison!D8"] == "=C8-B8"


def test_fill_xlsx_missing_named_input_and_approved_names(tmp_path: Path) -> None:
    import openpyxl

    payload = example_payload()
    del payload.texts["conclusion"]
    payload.items["mass_a"].quantity = Quantity.missing("kg")
    m = fill_xlsx(XLSX_TPL, payload, tmp_path / "c.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "c.xlsx")["Comparison"]
    assert ws["B21"].value == "MISSING" and ws["B8"].value == "MISSING"
    assert m.unknown_keys == ["conclusion"] and m.missing_keys == ["mass_a"]
    m2 = fill_xlsx(XLSX_TPL, example_payload(), tmp_path / "c2.xlsx", approved_names={"mass_a"})
    ws2 = openpyxl.load_workbook(tmp_path / "c2.xlsx")["Comparison"]
    assert ws2["B8"].value == 11.2 and ws2["C8"].value is None and "mass_b" in m2.unapproved_keys


def test_fill_xlsx_refuses_lossy_and_macro_templates(tmp_path: Path) -> None:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["a", 1])
    ws.append(["b", 2])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2))
    ws.add_chart(chart, "D2")
    wb.defined_names["mass_a"] = DefinedName("mass_a", attr_text=f"{ws.title}!$B$5")
    chart_tpl = tmp_path / "chart_tpl.xlsx"
    wb.save(chart_tpl)
    with pytest.raises(AgentError) as ei:
        fill_xlsx(chart_tpl, example_payload(), tmp_path / "out.xlsx")
    assert ei.value.code == "E_DOC_UNSUPPORTED" and "chart" in ei.value.details["lossy_elements"]
    m = fill_xlsx(chart_tpl, example_payload(), tmp_path / "out.xlsx", allow_lossy=True)
    assert any(u.element_type == "chart" and not u.preserved for u in m.unsupported_elements)
    assert openpyxl.load_workbook(tmp_path / "out.xlsx")[ws.title]["B5"].value == 11.2

    macro = tmp_path / "macro.xlsx"
    shutil.copyfile(XLSX_TPL, macro)
    with zipfile.ZipFile(macro, "a") as zf:
        zf.writestr("xl/vbaProject.bin", b"\x00")
    with pytest.raises(AgentError) as ei:
        fill_xlsx(macro, example_payload(), tmp_path / "m.xlsx")
    assert ei.value.code == "E_DOC_UNSUPPORTED" and not (tmp_path / "m.xlsx").exists()


def test_fill_xlsx_table_range_too_small_reports(tmp_path: Path) -> None:
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.load_workbook(XLSX_TPL)
    wb.defined_names["tbl_cost_lines"] = DefinedName("tbl_cost_lines", attr_text="Lines!$A$2:$B$3")
    tpl = tmp_path / "small.xlsx"
    wb.save(tpl)
    m = fill_xlsx(tpl, example_payload(), tmp_path / "small_out.xlsx")
    assert any(u.element_type == "table_range_too_small" for u in m.unsupported_elements)
    ws = openpyxl.load_workbook(tmp_path / "small_out.xlsx")["Lines"]
    assert ws["B3"].value == 780 and ws["B4"].value is None


def test_extra_payload_items_do_not_touch_template(tmp_path: Path) -> None:
    payload = example_payload()
    payload.items["unused"] = PayloadItem(
        key="unused", label_ko="미사용", quantity=Quantity(value=1.0, unit="", value_type="assumed", assumption=True)
    )
    m = fill_xlsx(XLSX_TPL, payload, tmp_path / "x.xlsx")
    assert "range:unused" not in m.changed_locators
    assert read_xlsx_structure(tmp_path / "x.xlsx")["defined_names"] == read_xlsx_structure(XLSX_TPL)["defined_names"]


def test_fill_pptx_refuses_macro_template_without_executing(tmp_path: Path) -> None:
    macro = tmp_path / "macro_template.pptx"
    shutil.copyfile(PPTX_TPL, macro)
    with zipfile.ZipFile(macro, "a") as zf:
        zf.writestr("ppt/vbaProject.bin", b"\x00fake-macro")
    with pytest.raises(AgentError) as ei:
        fill_pptx(macro, example_payload(), tmp_path / "out.pptx")
    assert ei.value.code == "E_DOC_UNSUPPORTED" and "macro" in ei.value.details["flags"]
    assert not (tmp_path / "out.pptx").exists()
