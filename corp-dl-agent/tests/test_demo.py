"""demo 통합 시나리오 시험 (BUILD_SPEC [16] J).

--fast 로 demo 전체(CAD A/B → 중량/원가 → 회귀/분류 run → predict → 근거 검색 → payload → PPTX/XLSX) 를 한 번 실행하고
필수 산출물 존재, demo_manifest 의 단계·검사 모두 PASS, 원본 fixture hash 불변, 문서 수치 == payload, outbound 0 을 확인한다.
한글/공백 경로 아래에서 실행한다. 네트워크 없음, 결정적(seed 고정).
"""

from __future__ import annotations

import csv
import json
import math
from pathlib import Path
from typing import Any

import pytest

from corp_dl_agent.cli import main
from corp_dl_agent.common import read_json, sha256_file
from corp_dl_agent.config import load_config
from corp_dl_agent.demo import (
    DEMO_MANIFEST_NAME,
    EXPECTED_MASS_KG,
    FIXTURE_FILES,
    DemoManifest,
    resolve_fixtures_dir,
    run_demo,
)

pytestmark = [pytest.mark.torch, pytest.mark.documents, pytest.mark.timeout(240)]

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures"

REQUIRED_OUTPUTS = (
    "candidates/A/cad_snapshot.json",
    "candidates/A/mass_result.json",
    "candidates/A/cost_breakdown.json",
    "candidates/A/candidate.json",
    "candidates/B/cad_snapshot.json",
    "candidates/B/mass_result.json",
    "candidates/B/cost_breakdown.json",
    "candidates/B/candidate.json",
    "design_comparison.json",
    "performance_predictions.csv",
    "performance_predictions.manifest.json",
    "retention_predictions.csv",
    "retention_predictions.manifest.json",
    "design_scenarios.csv",
    "design_scenario_predictions.csv",
    "design_scenario_predictions.manifest.json",
    "index_report.json",
    "evidence.json",
    "report_payload.json",
    "review.pptx",
    "comparison.xlsx",
    "document_manifest.json",
    "validation_report.json",
    DEMO_MANIFEST_NAME,
)


def _fixture_hashes() -> dict[str, str]:
    return {rel: sha256_file(FIXTURES / rel) for rel in FIXTURE_FILES.values()}


@pytest.fixture(scope="module")
def demo_run(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    """모듈당 1회: --fast 예산으로 demo 전체 실행 (한글/공백 경로)."""
    tmp = tmp_path_factory.mktemp("demo 시험")
    data_root = tmp / "작업 공간"
    out_dir = data_root / "outputs" / "demo 결과"
    before = _fixture_hashes()
    cfg = load_config(None, {"paths.data_root": str(data_root)})
    progress: list[str] = []
    manifest = run_demo(cfg, offline=True, device="cpu", out_dir=out_dir, fast=True, progress=progress.append)
    after = _fixture_hashes()
    return {
        "manifest": manifest,
        "out": out_dir,
        "data_root": data_root,
        "before": before,
        "after": after,
        "progress": progress,
    }


def test_demo_all_steps_and_checks_pass(demo_run: dict[str, Any]) -> None:
    m: DemoManifest = demo_run["manifest"]
    assert m.all_pass is True and m.exit_code == 0, [
        (s.name, s.status.value, s.detail) for s in m.steps if s.status.value != "PASS"
    ] + [(c.check_id, c.status.value, c.message_ko) for c in m.checks if c.status.value != "PASS"]
    assert [s.name for s in m.steps] == [
        "cad_import",
        "mass",
        "cost",
        "compare",
        "train_regression",
        "train_classification",
        "predict",
        "compare_performance",
        "docs_index",
        "docs_search",
        "payload",
        "docs_generate",
    ]
    assert all(s.status.value == "PASS" for s in m.steps)
    assert all(c.status.value == "PASS" for c in m.checks)
    assert m.synthetic is True and m.data_origin == "synthetic" and m.fast is True and m.device == "cpu"
    assert m.elapsed_seconds > 0 and m.finished_at
    # 단계 진행 메시지는 한국어로 출력된다
    assert any("중량 계산" in line for line in demo_run["progress"])
    # 파일에 저장된 manifest 도 동일하다
    saved = read_json(demo_run["out"] / DEMO_MANIFEST_NAME)
    assert (
        saved["all_pass"] is True
        and saved["exit_code"] == 0
        and saved["format"].startswith("corp-dl-agent/demo")
    )


def test_required_artifacts_exist_and_hashes_match(demo_run: dict[str, Any]) -> None:
    m: DemoManifest = demo_run["manifest"]
    out: Path = demo_run["out"]
    for rel in REQUIRED_OUTPUTS:
        assert (out / rel).is_file(), rel
    for rel, rec in m.artifacts.items():
        p = out / rel
        assert p.is_file(), rel
        assert sha256_file(p) == rec.sha256, rel
        assert rec.size == p.stat().st_size
    # run 산출물(summary/final_evaluation/export manifest) 도 산출물 폴더 아래(workspace) 에 있다
    run_artifacts = [rel for rel in m.artifacts if rel.startswith("workspace/runs/")]
    assert any(rel.endswith("export/manifest.json") for rel in run_artifacts)
    assert Path(m.workspace_dir).resolve() == (out / "workspace").resolve()


def test_original_fixtures_unchanged(demo_run: dict[str, Any]) -> None:
    m: DemoManifest = demo_run["manifest"]
    assert demo_run["before"] == demo_run["after"]
    assert m.fixtures and all(f.unchanged is True and f.sha256_after == f.sha256_before for f in m.fixtures)
    assert m.check("fixtures_unchanged").status.value == "PASS"
    assert m.check("templates_unchanged").status.value == "PASS"
    # fixture 폴더 안에 새 파일이 생기지 않았다 (색인 DB/run 은 산출물 폴더에만)
    assert not list((FIXTURES / "documents").glob("*.sqlite"))


def test_expected_masses_and_costs(demo_run: dict[str, Any]) -> None:
    out: Path = demo_run["out"]
    cmp = read_json(out / "design_comparison.json")
    rows = {r["candidate_id"]: r for r in cmp["rows"]}
    assert rows["A"]["mass_total"]["value"] == pytest.approx(EXPECTED_MASS_KG["A"])
    assert rows["B"]["mass_total"]["value"] == pytest.approx(EXPECTED_MASS_KG["B"])
    assert rows["A"]["unit_cost"]["value"] == pytest.approx(3782.8947368421054)
    assert rows["B"]["unit_cost"]["value"] == pytest.approx(3232.8947368421054)
    deltas = {d["metric"]: d for d in cmp["deltas"]}
    assert deltas["mass_total"]["delta"]["value"] == pytest.approx(3.5425 - 4.675)
    assert deltas["unit_cost"]["delta"]["value"] == pytest.approx(-550.0)
    assert cmp["consistent"] is True and cmp["synthetic"] is True and cmp["revision"] == "R1"
    # performance 는 predicted 출처로 비교에 포함된다
    perf = deltas["performance:insertion_force_n"]["delta"]
    assert perf["value"] is not None and perf["assumption"] is True
    for cand in ("A", "B"):
        pq = rows[cand]["performance"]["insertion_force_n"]
        assert pq["value_type"] == "predicted" and pq["model_run_id"] and pq["assumption"] is True
    mass_a = read_json(out / "candidates/A/mass_result.json")
    cube = next(i for i in mass_a["items"] if i["occurrence_path"] == "ROOT/ASM_A/PART_CUBE.1")
    assert cube["unit_mass"]["value"] == pytest.approx(1.0)


def test_predictions_marked_predicted_and_synthetic(demo_run: dict[str, Any]) -> None:
    m: DemoManifest = demo_run["manifest"]
    out: Path = demo_run["out"]
    run_ids = m.provenance["model_run_ids"]
    assert set(run_ids) == {"regression", "classification"}
    for name in ("performance_predictions", "retention_predictions", "design_scenario_predictions"):
        side = read_json(out / f"{name}.manifest.json")
        assert (
            side["value_type"] == "predicted"
            and side["synthetic"] is True
            and side["allow_synthetic"] is True
        )
        assert side["data_origin"] == "synthetic" and "demo" in side["demo_purpose"]
        assert side["n_rows"] > 0 and side["model_run_id"] in run_ids.values()
        with open(out / f"{name}.csv", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == side["n_rows"] and "prediction" in rows[0] and "in_train_range" in rows[0]
    scen = {
        r["id"]: r
        for r in csv.DictReader(open(out / "design_scenario_predictions.csv", encoding="utf-8-sig"))
    }
    assert set(scen) == {"DESIGN_A", "DESIGN_B"}
    payload = read_json(out / "report_payload.json")
    assert payload["synthetic"] is True and payload["data_origin"] == "synthetic"
    items = payload["items"]
    for key in ("insertion_force_a", "insertion_force_b", "retention_pass_rate"):
        q = items[key]["quantity"]
        assert q["value_type"] == "predicted" and q["model_run_id"] in run_ids.values(), key
    assert items["insertion_force_a"]["quantity"]["value"] == pytest.approx(
        float(scen["DESIGN_A"]["prediction"])
    )
    assert items["insertion_force_b"]["quantity"]["value"] == pytest.approx(
        float(scen["DESIGN_B"]["prediction"])
    )
    # scenario 입력의 두께는 CAD 파라미터에서 왔다 (값을 만들지 않음)
    thick = m.provenance["scenario"]["thickness"]
    assert thick["A"]["value_mm"] == pytest.approx(3.0) and thick["B"]["value_mm"] == pytest.approx(2.4)
    assert "param:nominal_thickness_mm" in thick["A"]["source_locator"]
    for key in ("mass_a", "mass_b", "mass_delta", "cost_a", "cost_b", "cost_delta", "cost_total"):
        assert items[key]["quantity"]["value_type"] == "calculated", key
    assert items["cost_total"]["quantity"]["value"] == pytest.approx(2200.0 + 157.89473684210526 + 500.0)


def test_document_numbers_match_payload(demo_run: dict[str, Any]) -> None:
    import openpyxl
    from pptx import Presentation

    out: Path = demo_run["out"]
    payload = read_json(out / "report_payload.json")
    items = payload["items"]
    report = read_json(out / "validation_report.json")
    assert report["passed"] is True and report["payload_passed"] is True
    assert all(d["n_fail"] == 0 for d in report["documents"])
    assert report["recalc_status"] == "RECALC_NOT_RUN" and report["render_status"] == "RENDER_NOT_RUN"
    # XLSX: named range 셀의 값 == payload
    wb = openpyxl.load_workbook(str(out / "comparison.xlsx"), data_only=False)
    try:
        for name in (
            "mass_a",
            "mass_b",
            "cost_a",
            "cost_b",
            "insertion_force_a",
            "insertion_force_b",
            "cost_total",
        ):
            dests = list(wb.defined_names[name].destinations)
            assert len(dests) == 1
            sheet, coord = dests[0]
            value = wb[sheet][coord.replace("$", "")].value
            assert isinstance(value, (int, float)) and math.isclose(
                float(value), float(items[name]["quantity"]["value"]), rel_tol=1e-9
            ), name
        sheet, coord = next(iter(wb.defined_names["cost_total_formula"].destinations))
        assert str(wb[sheet][coord.replace("$", "")].value).startswith("=SUM(")  # 수식 보존
        sheet, coord = next(iter(wb.defined_names["data_origin"].destinations))
        assert wb[sheet][coord.replace("$", "")].value == "synthetic"
    finally:
        wb.close()
    # PPTX: 렌더링된 값이 존재하고 placeholder/과거 값/MISSING 이 없다
    prs = Presentation(str(out / "review.pptx"))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text_frame.text)
    joined = "\n".join(texts)
    assert "4.675 kg" in joined and "3.5425 kg" in joined and "-1.1325" in joined
    assert "3,782.8947" in joined and "3,232.8947" in joined and "-550" in joined
    assert "{{" not in joined and "MISSING" not in joined and "CONFLICT" not in joined
    for stale in ("X-OLD", "2023-05-01", "1,234"):
        assert stale not in joined
    dm = read_json(out / "document_manifest.json")
    assert dm["synthetic"] is True and {d["document_type"] for d in dm["documents"]} == {"pptx", "xlsx"}
    assert all(d["template_hash_unchanged"] is True for d in dm["documents"])
    ev = read_json(out / "evidence.json")
    assert ev["n_items"] > 0 and all(i["source_id"] and i["locator"] and i["sha256"] for i in ev["items"])


def test_outbound_attempts_zero_and_guard_described(demo_run: dict[str, Any]) -> None:
    m: DemoManifest = demo_run["manifest"]
    assert (
        m.outbound["active"] is True and m.outbound["blocked_attempts"] == 0 and m.outbound["attempts"] == []
    )
    assert m.outbound["kind"] == "in_process_socket_guard" and m.outbound["os_level_isolation"] is False
    assert m.check("outbound_attempts_zero").status.value == "PASS"
    assert any("OS 수준" in n for n in m.notes)
    # manifest 에 개인 경로/실행 파일 경로가 없다
    env = m.environment
    assert "python_executable" not in env and "hostname" not in env and env["profile"] == "corp-offline"


def test_demo_cli_json_and_menu(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    data_root = tmp_path / "작업 공간"
    out = data_root / "outputs" / "demo"
    rc = main(
        [
            "demo",
            "--offline",
            "--device",
            "cpu",
            "--fast",
            "--output-dir",
            str(out),
            "--set",
            f"paths.data_root={data_root}",
            "--json",
        ]
    )
    captured = capsys.readouterr()
    assert rc == 0, captured.err[-2000:]
    payload = json.loads(captured.out)  # stdout 은 JSON 만 (진행 메시지는 stderr)
    assert payload["ok"] is True and payload["all_pass"] is True and payload["exit_code"] == 0
    assert (out / DEMO_MANIFEST_NAME).is_file() and "중량 계산" in captured.err
    # 메뉴 1번이 demo 로 노출된다
    from corp_dl_agent.commands.menu_cmd import available_items

    assert available_items()[0][0] == "1" and available_items()[0][3][0] == "demo"


def test_demo_cli_help_and_errors(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exc:
        main(["demo", "--help"])
    assert exc.value.code == 0
    out = capsys.readouterr().out
    assert "--fast" in out and "--fixtures-dir" in out and "오프라인" in out
    # 없는 fixture 폴더 → E_INPUT_INVALID (5)
    rc = main(
        [
            "demo",
            "--offline",
            "--device",
            "cpu",
            "--fast",
            "--fixtures-dir",
            str(tmp_path / "없는 폴더"),
            "--output-dir",
            str(tmp_path / "out"),
            "--set",
            f"paths.data_root={tmp_path / 'ws'}",
            "--json",
        ]
    )
    err = json.loads(capsys.readouterr().out)
    assert rc == 5 and err["ok"] is False and err["error"]["code"] == "E_INPUT_INVALID"
    # 필수 fixture 파일이 빠진 폴더 → 누락 목록 보고
    partial = tmp_path / "부분 fixture"
    (partial / "cad").mkdir(parents=True)
    rc = main(
        [
            "demo",
            "--fast",
            "--fixtures-dir",
            str(partial),
            "--output-dir",
            str(tmp_path / "out2"),
            "--set",
            f"paths.data_root={tmp_path / 'ws'}",
            "--json",
        ]
    )
    err = json.loads(capsys.readouterr().out)
    assert (
        rc == 5
        and "missing" in err["error"]["details"]
        and "ml/task_regression.yaml" in err["error"]["details"]["missing"]
    )
    # 승인 root 밖 산출물 폴더 → E_PATH_OUTSIDE_ROOT
    cfg = load_config(None, {"paths.data_root": str(tmp_path / "ws")})
    from corp_dl_agent.demo import resolve_output_dir
    from corp_dl_agent.errors import AgentError

    with pytest.raises(AgentError) as ae:
        resolve_output_dir(cfg, tmp_path / "밖" / "demo")
    assert ae.value.code == "E_PATH_OUTSIDE_ROOT"


def test_demo_device_cuda_not_supported_without_gpu(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    import torch

    if torch.cuda.is_available():
        pytest.skip("CUDA 가 사용 가능한 환경에서는 이 거부 경로를 시험하지 않는다")
    data_root = tmp_path / "ws"
    rc = main(
        [
            "demo",
            "--device",
            "cuda",
            "--fast",
            "--output-dir",
            str(data_root / "outputs" / "demo"),
            "--set",
            f"paths.data_root={data_root}",
            "--json",
        ]
    )
    err = json.loads(capsys.readouterr().out)
    assert (
        rc == 9
        and err["error"]["code"] == "E_NOT_SUPPORTED"
        and err["error"]["details"]["status"] == "NOT_RUN"
    )
    assert not (data_root / "outputs" / "demo").exists()  # 장치 검사는 산출물 폴더를 만들기 전에 수행된다


def test_resolve_fixtures_dir_env_and_project(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from corp_dl_agent.errors import AgentError

    assert resolve_fixtures_dir(None) == FIXTURES.resolve()
    assert resolve_fixtures_dir(FIXTURES) == FIXTURES.resolve()
    monkeypatch.setenv("DIA_FIXTURES_DIR", str(tmp_path / "없음"))
    with pytest.raises(AgentError) as exc:
        resolve_fixtures_dir(None)
    assert exc.value.code == "E_INPUT_INVALID"
    monkeypatch.setenv("DIA_FIXTURES_DIR", str(FIXTURES))
    assert resolve_fixtures_dir(None) == FIXTURES.resolve()
