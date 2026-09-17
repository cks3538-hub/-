"""documents.validation: 수치 일치·구조 보존·오래된 값 잔존 회귀·hash·RECALC/RENDER 상태."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from corp_dl_agent.common import Quantity, Status
from corp_dl_agent.documents.payload import validate_payload
from corp_dl_agent.documents.synthetic import example_payload
from corp_dl_agent.documents.templates import DocumentManifest, collect_pptx_texts, collect_xlsx_texts, fill_pptx, fill_xlsx
from corp_dl_agent.documents.validation import (
    ValidationReport,
    find_stale,
    stale_pattern,
    validate_documents,
    write_validation_report,
)

pytestmark = pytest.mark.documents

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "documents"


@pytest.fixture
def generated(tmp_path: Path) -> tuple[Path, Path, list[DocumentManifest]]:
    tpl_dir = tmp_path / "템플릿"
    shutil.copytree(FIXTURES, tpl_dir)
    out = tmp_path / "산출물"
    payload = validate_payload(example_payload()).payload
    m1 = fill_pptx(tpl_dir / "template_review.pptx", payload, out / "review.pptx")
    m2 = fill_xlsx(tpl_dir / "template_comparison.xlsx", payload, out / "comparison.xlsx")
    return out / "review.pptx", out / "comparison.xlsx", [m1, m2]


def _check(report: ValidationReport, doc_type: str, check_id: str):  # type: ignore[no-untyped-def]
    doc = next(d for d in report.documents if d.document_type == doc_type)
    return next(c for c in doc.checks if c.check_id == check_id)


def test_full_validation_passes_and_reports_statuses(generated, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    pptx, xlsx, manifests = generated
    report = validate_documents(pptx, xlsx, example_payload(), manifests=manifests)
    assert report.passed and report.payload_passed and report.synthetic and report.n_conflict == 0
    assert report.recalc_status == "RECALC_NOT_RUN" and report.render_status == "RENDER_NOT_RUN"
    assert _check(report, "pptx", "pptx.numeric").status is Status.PASS
    assert _check(report, "pptx", "pptx.structure").status is Status.PASS
    assert _check(report, "pptx", "pptx.template_hash").status is Status.PASS
    assert _check(report, "pptx", "pptx.stale_values").status is Status.PASS
    assert _check(report, "pptx", "pptx.placeholders").status is Status.PASS
    assert _check(report, "pptx", "pptx.render").status is Status.NOT_RUN
    assert _check(report, "xlsx", "xlsx.numeric").status is Status.PASS
    assert _check(report, "xlsx", "xlsx.structure").status is Status.PASS
    assert _check(report, "xlsx", "xlsx.recalc").status is Status.NOT_RUN
    assert _check(report, "xlsx", "xlsx.stale_values").status is Status.PASS
    assert report.stale_values == ["X-OLD", "2023-05-01", "1,234"]
    xlsx_doc = next(d for d in report.documents if d.document_type == "xlsx")
    assert xlsx_doc.n_fail == 0 and xlsx_doc.n_not_run == 2 and xlsx_doc.passed
    p = write_validation_report(report, tmp_path / "validation_report.json")
    assert p.is_file() and "RECALC_NOT_RUN" in p.read_text(encoding="utf-8")


def test_regression_new_values_present_and_old_values_absent(generated) -> None:  # type: ignore[no-untyped-def]
    pptx, xlsx, _ = generated
    pptx_text = "\n".join(collect_pptx_texts(pptx).values())
    xlsx_text = "\n".join(collect_xlsx_texts(xlsx).values())
    for new in ("X-NEW", "2026-09-01", "1,980"):
        assert new in pptx_text
    assert "X-NEW" in xlsx_text and "2026-09-01" in xlsx_text and "1980" in xlsx_text
    for old in ("X-OLD", "2023-05-01", "1,234", "1234"):
        assert old not in pptx_text and old not in xlsx_text
    assert not find_stale(collect_pptx_texts(pptx), ["X-OLD", "2023-05-01", "1,234"])


def test_stale_value_injected_into_payload_fails_validation(tmp_path: Path) -> None:
    payload = example_payload()
    payload.texts["conclusion"].text = "X-OLD 대비 유리 (구 단가 1,234 원 기준)"
    out = tmp_path / "out"
    m1 = fill_pptx(FIXTURES / "template_review.pptx", payload, out / "review.pptx")
    m2 = fill_xlsx(FIXTURES / "template_comparison.xlsx", payload, out / "comparison.xlsx")
    report = validate_documents(out / "review.pptx", out / "comparison.xlsx", payload, manifests=[m1, m2])
    assert not report.passed
    c = _check(report, "pptx", "pptx.stale_values")
    assert c.status is Status.FAIL and "slide:2/shape:Body 1" in (c.locator or "") and "X-OLD" in (c.actual or "")
    assert _check(report, "xlsx", "xlsx.stale_values").status is Status.FAIL


def test_stale_pattern_numeric_boundaries() -> None:
    pat = stale_pattern("1,234")
    assert pat.search("단가 1,234 원") and pat.search("단가 1234원") and not pat.search("11234") and not pat.search("1,2345")
    assert stale_pattern("X-OLD").search("차종 x-old-b")
    assert find_stale({"a": "값 12.5"}, ["12.5"]) == [("a", "12.5")]
    assert find_stale({"a": "값 112.5"}, ["12.5"]) == []


def test_tampered_output_and_template_fail_hash_and_numeric(generated, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    import openpyxl

    pptx, xlsx, manifests = generated
    wb = openpyxl.load_workbook(xlsx)
    wb["Comparison"]["B8"] = 99.0
    wb.save(xlsx)
    report = validate_documents(pptx, xlsx, example_payload(), manifests=manifests)
    assert not report.passed
    assert _check(report, "xlsx", "xlsx.output_hash").status is Status.FAIL
    assert _check(report, "xlsx", "xlsx.numeric").status is Status.FAIL
    assert _check(report, "xlsx", "xlsx.numeric[range:mass_a]").status is Status.FAIL
    assert _check(report, "pptx", "pptx.output_hash").status is Status.PASS
    # 템플릿 원본 변경 → template_hash FAIL
    tpl = Path(manifests[0].template_path)
    tpl.write_bytes(tpl.read_bytes() + b"\x00")
    report2 = validate_documents(pptx, None, example_payload(), manifests=manifests)
    assert _check(report2, "pptx", "pptx.template_hash").status is Status.FAIL


def test_structure_change_detected(generated) -> None:  # type: ignore[no-untyped-def]
    from pptx import Presentation

    pptx, xlsx, manifests = generated
    prs = Presentation(str(pptx))
    sld_id_lst = prs.slides._sldIdLst
    sld_id_lst.remove(sld_id_lst[-1])
    prs.save(str(pptx))
    report = validate_documents(pptx, None, example_payload(), manifests=manifests)
    c = _check(report, "pptx", "pptx.structure")
    assert c.status is Status.FAIL and "n_slides" in c.message_ko
    import openpyxl

    wb = openpyxl.load_workbook(xlsx)
    wb["Comparison"]["D8"] = "=B8-C8"
    del wb.defined_names["mass_b"]
    wb.save(xlsx)
    report2 = validate_documents(None, xlsx, example_payload(), manifests=manifests)
    c2 = _check(report2, "xlsx", "xlsx.structure")
    assert c2.status is Status.FAIL and "defined_names" in c2.message_ko and "formulas" in c2.message_ko


def test_missing_values_are_partial_not_fail_and_conflict_fails(tmp_path: Path) -> None:
    payload = example_payload()
    payload.items["mass_b"].quantity = Quantity.missing("kg")
    out = tmp_path / "out"
    m1 = fill_pptx(FIXTURES / "template_review.pptx", payload, out / "review.pptx")
    m2 = fill_xlsx(FIXTURES / "template_comparison.xlsx", payload, out / "comparison.xlsx")
    report = validate_documents(out / "review.pptx", out / "comparison.xlsx", payload, manifests=[m1, m2])
    # mass_delta 는 mass_b 가 없어 재계산 불가 → payload 검증 실패(CONFLICT) 가 보고에 남는다
    assert report.n_missing >= 1 and report.n_conflict == 1 and not report.payload_passed and not report.passed
    pptx_doc = next(d for d in report.documents if d.document_type == "pptx")
    assert any(c.status is Status.PARTIAL and c.check_id.startswith("pptx.missing") for c in pptx_doc.checks)
    assert any(c.status is Status.FAIL and c.check_id.startswith("pptx.conflict") for c in pptx_doc.checks)


def test_validation_without_manifests_marks_not_run(generated) -> None:  # type: ignore[no-untyped-def]
    pptx, xlsx, _ = generated
    report = validate_documents(pptx, xlsx, example_payload())
    assert _check(report, "pptx", "pptx.numeric").status is Status.NOT_RUN
    assert _check(report, "pptx", "pptx.structure").status is Status.NOT_RUN
    assert _check(report, "xlsx", "xlsx.template_hash").status is Status.NOT_RUN
    assert _check(report, "pptx", "pptx.stale_values").status is Status.PASS
    assert report.passed


def test_xlsx_conflict_after_revalidation_and_unbacked_value(tmp_path: Path) -> None:
    """manifest 는 값이 채워진 payload 로 만들고, 검증은 mass_a 가 사라진 payload 로 한다.

    - mass_a: 문서에는 11.2 가 있으나 payload 에 값이 없음 → FAIL (근거 없는 숫자)
    - mass_delta: mass_a 누락으로 재계산 불가 → CONFLICT → xlsx.conflict FAIL
    """
    filled = example_payload()
    out = tmp_path / "out"
    m1 = fill_pptx(FIXTURES / "template_review.pptx", filled, out / "review.pptx")
    m2 = fill_xlsx(FIXTURES / "template_comparison.xlsx", filled, out / "comparison.xlsx")
    later = example_payload()
    later.items["mass_a"].quantity = Quantity.missing("kg", "재검토에서 제외됨")
    report = validate_documents(out / "review.pptx", out / "comparison.xlsx", later, manifests=[m1, m2])
    assert not report.passed and report.n_conflict == 1 and report.n_missing == 2
    xlsx_doc = next(d for d in report.documents if d.document_type == "xlsx")
    ids = {c.check_id: c for c in xlsx_doc.checks}
    assert ids["xlsx.numeric[range:mass_a]"].status is Status.FAIL and ids["xlsx.numeric[range:mass_a]"].expected == "MISSING"
    assert ids["xlsx.numeric[range:mass_a]"].actual == "11.2"
    conflict = next(c for k, c in ids.items() if k.startswith("xlsx.conflict[") and "mass_delta" in k)
    assert conflict.status is Status.FAIL and conflict.expected == "CONFLICT" and "mass_delta_diff" in conflict.message_ko
    assert ids["xlsx.numeric"].status is Status.FAIL
    # 표 셀(tbl_cost_lines) 값은 그대로 일치한다
    assert not any(k.startswith("xlsx.numeric[range:tbl_cost_lines") for k in ids)
    pptx_doc = next(d for d in report.documents if d.document_type == "pptx")
    assert any(c.check_id.startswith("pptx.conflict[") and c.status is Status.FAIL for c in pptx_doc.checks)
    assert any(c.check_id == "pptx.numeric[slide:2/shape:Body 1:mass_a]" and c.expected == "MISSING" for c in pptx_doc.checks)


def test_engine_configured_but_unsupported_stays_not_run(generated) -> None:  # type: ignore[no-untyped-def]
    pptx, xlsx, manifests = generated
    report = validate_documents(
        pptx, xlsx, example_payload(), manifests=manifests, renderer="libreoffice", recalc_engine="excel_com"
    )
    assert report.passed and report.renderer == "libreoffice" and report.recalc_engine == "excel_com"
    assert report.recalc_status == "RECALC_NOT_RUN" and report.render_status == "RENDER_NOT_RUN"
    for doc_type, check_id in (("pptx", "pptx.render"), ("xlsx", "xlsx.render"), ("xlsx", "xlsx.recalc")):
        c = _check(report, doc_type, check_id)
        assert c.status is Status.NOT_RUN and "E_NOT_SUPPORTED" in c.message_ko
    assert _check(report, "xlsx", "xlsx.recalc").actual == "formula_cells=6, recalc_engine=excel_com"
    assert sum("E_NOT_SUPPORTED" in n for n in report.notes) == 2
    default = validate_documents(pptx, xlsx, example_payload(), manifests=manifests)
    assert "E_NOT_SUPPORTED" not in _check(default, "pptx", "pptx.render").message_ko
    assert not any("E_NOT_SUPPORTED" in n for n in default.notes)
