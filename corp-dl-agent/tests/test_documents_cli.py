"""CLI 시험: docs index / search / generate / make-synthetic 가 실제 산출물을 만든다 (한글 경로 포함)."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest

from corp_dl_agent.cli import main
from corp_dl_agent.common import read_json, sha256_file
from corp_dl_agent.errors import EXIT_NOT_SUPPORTED, EXIT_USAGE, EXIT_VALIDATION

pytestmark = pytest.mark.documents

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "documents"


@pytest.fixture
def env(tmp_path: Path) -> dict[str, Path]:
    src = tmp_path / "근거 문서"
    shutil.copytree(FIXTURES, src)
    ws = tmp_path / "작업 공간"
    return {"src": src, "ws": ws, "tmp": tmp_path}


def _common(env: dict[str, Path]) -> list[str]:
    return ["--set", f"paths.data_root={env['ws']}", "--set", f"paths.sources_roots=[{env['src']}]"]


def _json_out(capsys: pytest.CaptureFixture[str]) -> dict:  # type: ignore[type-arg]
    out = capsys.readouterr().out
    return json.loads(out)  # type: ignore[no-any-return]


def test_docs_index_and_search_json(env: dict[str, Path], capsys: pytest.CaptureFixture[str]) -> None:
    rc = main(["docs", "index", "--json", *_common(env)])
    assert rc == 0
    data = _json_out(capsys)
    assert data["ok"] and data["stats"]["n_documents"] == 5 and data["stats"]["n_hidden_fragments"] == 5
    assert Path(data["db"]) == env["ws"] / "indexes" / "documents.sqlite" and Path(data["db"]).is_file()
    assert len(data["report"]["needs_conversion"]) == 1

    rc = main(["docs", "search", "--query", "1,234 원", "--json", *_common(env)])
    assert rc == 0
    data = _json_out(capsys)
    assert data["n_hits"] > 0 and all(not h["hidden"] for h in data["hits"])
    locs = {h["locator"] for h in data["hits"]}
    assert "slide:2/shape:Body 1" in locs and "sheet:Cost!B11" in locs

    rc = main(["docs", "search", "--query", "1,100", "--json", *_common(env)])
    assert _json_out(capsys)["n_hits"] == 0
    rc = main(["docs", "search", "--query", "1,100", "--include-hidden", "--json", *_common(env)])
    assert rc == 0 and _json_out(capsys)["n_hits"] == 2

    ev_path = env["ws"] / "outputs" / "evidence.json"
    rc = main(["docs", "search", "--query", "X-OLD 단가", "--evidence-output", str(ev_path), *_common(env)])
    assert rc == 0 and ev_path.is_file()
    text = capsys.readouterr().out
    assert "검색어" in text and "evidence 저장" in text
    ev = read_json(ev_path)
    assert ev["n_items"] > 0 and ev["items"][0]["sha256"] and ev["items"][0]["kind"] == "fact"


def test_docs_index_rejects_root_outside_sources(
    env: dict[str, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    outside = env["tmp"] / "밖"
    outside.mkdir()
    rc = main(["docs", "index", "--roots", str(outside), *_common(env)])
    assert rc == EXIT_VALIDATION
    assert "E_PATH_OUTSIDE_ROOT" in capsys.readouterr().err
    # sources_roots 가 비어 있고 --roots 도 없으면 사용법 오류
    rc = main(["docs", "index", "--set", f"paths.data_root={env['ws']}"])
    assert rc == EXIT_USAGE


def test_docs_generate_creates_all_outputs(env: dict[str, Path], capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["docs", "index", *_common(env)]) == 0
    capsys.readouterr()
    out_dir = env["ws"] / "outputs" / "검토 보고"
    tpl_pptx = env["src"] / "template_review.pptx"
    tpl_xlsx = env["src"] / "template_comparison.xlsx"
    h_pptx, h_xlsx = sha256_file(tpl_pptx), sha256_file(tpl_xlsx)
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(env["src"] / "report_payload.example.json"),
            "--pptx-template",
            str(tpl_pptx),
            "--xlsx-template",
            str(tpl_xlsx),
            "--output-dir",
            str(out_dir),
            "--query",
            "X-OLD 단가",
            "--query",
            "중량",
            "--json",
            *_common(env),
        ]
    )
    assert rc == 0
    data = _json_out(capsys)
    assert data["ok"] is True and data["payload_validation"]["passed"] is True
    for name in (
        "review.pptx",
        "comparison.xlsx",
        "evidence.json",
        "report_payload.json",
        "document_manifest.json",
        "validation_report.json",
    ):
        assert (out_dir / name).is_file(), name
    assert sha256_file(tpl_pptx) == h_pptx and sha256_file(tpl_xlsx) == h_xlsx
    manifest = read_json(out_dir / "document_manifest.json")
    assert manifest["synthetic"] is True and manifest["recalc_status"] == "RECALC_NOT_RUN"
    assert {d["document_type"] for d in manifest["documents"]} == {"pptx", "xlsx"}
    assert all(d["template_hash_unchanged"] for d in manifest["documents"])
    assert manifest["payload_hash"] == sha256_file(out_dir / "report_payload.json")
    report = read_json(out_dir / "validation_report.json")
    assert report["passed"] is True and report["render_status"] == "RENDER_NOT_RUN"
    assert report["stale_values"] == ["X-OLD", "2023-05-01", "1,234"]
    ev = read_json(out_dir / "evidence.json")
    assert ev["n_items"] > 0 and ev["synthetic"] is True and ev["excluded_hidden"] == 0
    assert all(i["kind"] == "fact" for i in ev["items"]) is False or ev["n_form_reference"] == 0
    payload_out = read_json(out_dir / "report_payload.json")
    assert payload_out["items"]["cost_total"]["quantity"]["value"] == 1980


def test_docs_generate_text_output_and_evidence_file(
    env: dict[str, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    out_dir = env["ws"] / "outputs" / "gen2"
    ev_in = env["ws"] / "outputs" / "prev_evidence.json"
    ev_in.parent.mkdir(parents=True)
    from corp_dl_agent.documents.evidence import build_evidence, write_evidence

    write_evidence(build_evidence([], purpose="이전 검색"), ev_in)
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(env["src"] / "report_payload.example.json"),
            "--pptx-template",
            str(env["src"] / "template_review.pptx"),
            "--output-dir",
            str(out_dir),
            "--evidence",
            str(ev_in),
            *_common(env),
        ]
    )
    text = capsys.readouterr().out
    assert rc == 0 and "산출물 폴더" in text and "review.pptx" in text and "RECALC_NOT_RUN" in text
    assert (out_dir / "review.pptx").is_file() and not (out_dir / "comparison.xlsx").exists()
    assert read_json(out_dir / "evidence.json")["purpose"] == "이전 검색"


def test_docs_generate_fails_validation_exit_code(
    env: dict[str, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    out_dir = env["ws"] / "outputs" / "gen3"
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(env["src"] / "report_payload.example.json"),
            "--xlsx-template",
            str(env["src"] / "template_comparison.xlsx"),
            "--output-dir",
            str(out_dir),
            "--stale",
            "X-NEW",
            "--json",
            *_common(env),
        ]
    )
    assert rc == EXIT_VALIDATION
    data = _json_out(capsys)
    assert data["ok"] is False and (out_dir / "validation_report.json").is_file()
    assert any(
        c["check_id"] == "xlsx.stale_values" and c["status"] == "FAIL"
        for c in data["validation"]["documents"][0]["checks"]
    )


def test_docs_generate_usage_and_unsupported(
    env: dict[str, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(env["src"] / "report_payload.example.json"),
            "--output-dir",
            str(env["ws"] / "outputs" / "x"),
            *_common(env),
        ]
    )
    assert rc == EXIT_USAGE
    import zipfile

    macro = env["src"] / "macro_template.xlsx"
    shutil.copyfile(FIXTURES / "template_comparison.xlsx", macro)
    with zipfile.ZipFile(macro, "a") as zf:
        zf.writestr("xl/vbaProject.bin", b"\x00")
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(env["src"] / "report_payload.example.json"),
            "--xlsx-template",
            str(macro),
            "--output-dir",
            str(env["ws"] / "outputs" / "y"),
            *_common(env),
        ]
    )
    assert rc == EXIT_NOT_SUPPORTED and "E_DOC_UNSUPPORTED" in capsys.readouterr().err


def test_docs_make_synthetic(env: dict[str, Path], capsys: pytest.CaptureFixture[str]) -> None:
    out = env["ws"] / "outputs" / "합성"
    rc = main(["docs", "make-synthetic", "--output-dir", str(out), "--json", *_common(env)])
    assert rc == 0
    data = _json_out(capsys)
    assert (
        data["synthetic"] is True
        and (out / "README.md").is_file()
        and (out / "past_review_2023.pptx").is_file()
    )
    assert "synthetic" in (out / "README.md").read_text(encoding="utf-8")


def test_docs_help_lists_subcommands(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as ei:
        main(["docs", "--help"])
    assert ei.value.code == 0
    out = capsys.readouterr().out
    assert "index" in out and "search" in out and "generate" in out and "색인" in out


def test_docs_generate_wrong_sum_becomes_conflict_not_fabricated(
    env: dict[str, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    """payload 의 합계가 항목 합과 다르면 계산 엔진 재검증이 CONFLICT 로 바꾸고 문서에는 CONFLICT 가 기록된다."""
    import openpyxl

    data = read_json(FIXTURES / "report_payload.example.json")
    data["items"]["cost_total"]["quantity"]["value"] = 9999
    bad_payload = env["src"] / "payload_wrong_sum.json"
    bad_payload.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    out_dir = env["ws"] / "outputs" / "상충 검토"
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(bad_payload),
            "--pptx-template",
            str(env["src"] / "template_review.pptx"),
            "--xlsx-template",
            str(env["src"] / "template_comparison.xlsx"),
            "--output-dir",
            str(out_dir),
            "--json",
            *_common(env),
        ]
    )
    assert rc == EXIT_VALIDATION
    data_out = _json_out(capsys)
    assert data_out["ok"] is False and data_out["payload_validation"]["passed"] is False
    assert data_out["payload_validation"]["n_conflict"] == 1
    by_id = {c["check_id"]: c for c in data_out["payload_validation"]["checks"]}
    assert by_id["cost_total_sum"]["status"] == "FAIL" and by_id["cost_total_sum"]["expected"] == 1980
    assert by_id["cost_total_sum"]["actual"] == 9999
    # 같은 대상의 두 번째 검사(표 합)는 CONFLICT 를 계산값으로 덮어쓰지 않고 NOT_RUN 으로 남는다
    assert (
        by_id["cost_total_table_sum"]["status"] == "NOT_RUN"
        and by_id["cost_total_table_sum"]["expected"] == 1980
    )
    written = read_json(out_dir / "report_payload.json")
    q = written["items"]["cost_total"]["quantity"]
    assert (
        q["value_type"] == "conflict"
        and q["value"] is None
        and "9,999" in q["notes"]
        and "1,980" in q["notes"]
    )
    ws = openpyxl.load_workbook(out_dir / "comparison.xlsx")["Comparison"]
    assert ws["B17"].value == "CONFLICT" and ws["B16"].value == "=SUM(B13:B15)"
    from corp_dl_agent.documents.templates import collect_pptx_texts

    body = collect_pptx_texts(out_dir / "review.pptx")["slide:2/shape:Body 1"]
    assert "원가 합계: CONFLICT" in body and "9,999" not in body and "9999" not in body
    report = read_json(out_dir / "validation_report.json")
    assert report["passed"] is False and report["n_conflict"] == 1
    xlsx_doc = next(d for d in report["documents"] if d["document_type"] == "xlsx")
    assert any(
        c["check_id"] == "xlsx.conflict[range:cost_total]" and c["status"] == "FAIL"
        for c in xlsx_doc["checks"]
    )
    manifest = read_json(out_dir / "document_manifest.json")
    assert all("cost_total" in d["missing_keys"] for d in manifest["documents"])


def test_docs_generate_reports_unsupported_engines_as_not_run(
    env: dict[str, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    out_dir = env["ws"] / "outputs" / "엔진 설정"
    rc = main(
        [
            "docs",
            "generate",
            "--payload",
            str(env["src"] / "report_payload.example.json"),
            "--xlsx-template",
            str(env["src"] / "template_comparison.xlsx"),
            "--output-dir",
            str(out_dir),
            "--set",
            "documents.renderer=libreoffice",
            "--set",
            "documents.recalc_engine=excel_com",
            "--json",
            *_common(env),
        ]
    )
    assert rc == 0
    data = _json_out(capsys)
    assert data["ok"] is True
    report = read_json(out_dir / "validation_report.json")
    assert report["renderer"] == "libreoffice" and report["recalc_engine"] == "excel_com"
    assert report["recalc_status"] == "RECALC_NOT_RUN" and report["render_status"] == "RENDER_NOT_RUN"
    assert sum("E_NOT_SUPPORTED" in n for n in report["notes"]) == 2
    statuses = {c["check_id"]: c for c in report["documents"][0]["checks"]}
    assert (
        statuses["xlsx.recalc"]["status"] == "NOT_RUN"
        and "excel_com" in statuses["xlsx.recalc"]["message_ko"]
    )
    assert (
        statuses["xlsx.render"]["status"] == "NOT_RUN"
        and "libreoffice" in statuses["xlsx.render"]["message_ko"]
    )
    manifest = read_json(out_dir / "document_manifest.json")
    assert "renderer=libreoffice" in " ".join(manifest["notes"])
