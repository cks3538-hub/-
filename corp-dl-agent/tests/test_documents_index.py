"""documents.index: PPTX/XLSX 색인 — locator, 수식/cached 구분, hidden, revision, prune, 구형 형식, root 제한."""

from __future__ import annotations

import os
import shutil
import zipfile
from pathlib import Path

import pytest

from corp_dl_agent.documents.index import (
    DocumentIndex,
    detect_date,
    detect_unit,
    extract_pptx,
    extract_xlsx,
    normalize_approval,
    revision_from_filename,
    scan_package_flags,
)
from corp_dl_agent.errors import AgentError

pytestmark = pytest.mark.documents

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "documents"


@pytest.fixture
def docs_dir(tmp_path: Path) -> Path:
    d = tmp_path / "문서 자료"
    shutil.copytree(FIXTURES, d)
    return d


def _by_name(index: DocumentIndex, name: str):  # type: ignore[no-untyped-def]
    for d in index.documents():
        if Path(d.path).name == name:
            return d
    raise AssertionError(f"not indexed: {name}")


def test_index_roots_reports_statuses_and_flags(docs_dir: Path, tmp_path: Path) -> None:
    with DocumentIndex(tmp_path / "idx" / "documents.sqlite") as index:
        rep = index.index_roots([docs_dir])
        assert len(rep.indexed) == 4 and rep.unchanged == [] and rep.errors == []
        assert [Path(p).name for p in rep.needs_conversion] == ["legacy_note.ppt"]
        assert rep.n_documents == 5 and rep.n_hidden_fragments == 5
        legacy = _by_name(index, "legacy_note.ppt")
        assert legacy.status == "NEEDS_APPROVED_CONVERSION" and legacy.n_fragments == 0 and legacy.sha256
        past = _by_name(index, "past_review_2023.pptx")
        assert past.status == "INDEXED" and past.kind == "pptx" and past.synthetic is True
        assert past.revision == "3" and past.approval_state == "approved" and past.doc_kind == "fact"
        assert "chart" in past.flags and "macro" not in past.flags
        cost = _by_name(index, "past_cost_table.xlsx")
        assert cost.revision == "2" and cost.approval_state == "approved" and cost.hidden_fragments == 3
        tpl = _by_name(index, "template_review.pptx")
        assert tpl.doc_kind == "form_reference" and tpl.approval_state == "draft"
        assert _by_name(index, "template_comparison.xlsx").doc_kind == "form_reference"
        # 재색인: 변경 없음
        rep2 = index.index_roots([docs_dir])
        assert rep2.indexed == [] and len(rep2.unchanged) == 4 and rep2.pruned == []


def test_pptx_fragments_locators_hidden_chart_notes(docs_dir: Path) -> None:
    frags, props = extract_pptx(docs_dir / "past_review_2023.pptx")
    by_loc = {f.locator: f for f in frags}
    assert by_loc["slide:1/shape:Title 1"].text == "X-OLD 설계 검토서"
    assert by_loc["slide:2/table:Table 1/r2c2"].text == "1,234"
    assert by_loc["slide:2/table:Table 1/r2c3"].text == "원"
    assert "2023-05-01" in by_loc["slide:2/notes"].text and by_loc["slide:2/notes"].fragment_kind == "notes"
    body = by_loc["slide:2/shape:Body 1"]
    assert body.unit == "원" and body.base_date == "2023-05-01"
    chart = by_loc["slide:3/chart:Chart 1/series:중량(kg)"]
    assert chart.fragment_kind == "chart_series" and chart.text == "중량(kg): X-OLD=12.5; X-OLD-B=11.8"
    assert chart.meta["categories"] == '["X-OLD", "X-OLD-B"]'
    assert by_loc["slide:3/chart:Chart 1/title"].text == "중량 비교"
    assert by_loc["slide:4/shape:Memo 1"].hidden is True and by_loc["slide:4/shape:Title 1"].hidden is True
    assert not by_loc["slide:1/shape:Title 1"].hidden
    assert props["n_slides"] == 4 and props["hidden_slides"] == 1 and props["category"] == "synthetic"


def test_xlsx_fragments_formula_vs_cached_named_range_table_hidden(docs_dir: Path) -> None:
    frags, props = extract_xlsx(docs_dir / "past_cost_table.xlsx")
    by_loc = {f.locator: f for f in frags}
    total = by_loc["sheet:Cost!B7"]
    assert total.formula == "=SUM(B2:B6)" and total.cached_value == "1234" and total.number_format == "#,##0"
    stale = by_loc["sheet:Cost!B9"]
    assert stale.formula == "=B7*0.1" and stale.cached_value == "120" and stale.hidden is True
    assert by_loc["sheet:Cost!B2"].formula is None and by_loc["sheet:Cost!B2"].text == "500"
    assert by_loc["range:UnitCost"].formula == "=SUM(B2:B6)" and by_loc["range:UnitCost"].cached_value == "1234"
    assert by_loc["range:BaseDate"].base_date == "2023-05-01"
    assert by_loc["range:Vehicle"].text == "X-OLD"
    assert by_loc["table:T1!r1c2"].text == "금액(원)" and by_loc["table:T1!r2c2"].text == "500"
    assert by_loc["sheet:내부메모!A1"].hidden is True
    assert by_loc["sheet:Cost!B11"].unit == "원"
    assert props["hidden_sheets"] == ["내부메모"] and props["sheets"] == ["Cost", "내부메모"]


def test_hidden_content_not_stored_when_flag_off(docs_dir: Path, tmp_path: Path) -> None:
    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        rep = index.index_roots([docs_dir], include_hidden_flag=False)
        assert rep.n_hidden_fragments == 0
        past = _by_name(index, "past_review_2023.pptx")
        locs = {f.locator for f in index.fragments(past.source_id)}
        assert "slide:4/shape:Memo 1" not in locs and "slide:2/shape:Body 1" in locs


def test_prune_removes_deleted_and_reindexes_modified(docs_dir: Path, tmp_path: Path) -> None:
    import openpyxl

    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        index.index_roots([docs_dir])
        cost = _by_name(index, "past_cost_table.xlsx")
        old_hash = cost.sha256
        (docs_dir / "legacy_note.ppt").unlink()
        wb = openpyxl.load_workbook(docs_dir / "past_cost_table.xlsx")
        wb["Cost"]["C2"] = "변경됨"
        wb.save(docs_dir / "past_cost_table.xlsx")
        rep = index.index_roots([docs_dir])
        assert [Path(p).name for p in rep.pruned] == ["legacy_note.ppt"]
        assert [Path(p).name for p in rep.indexed] == ["past_cost_table.xlsx"]
        assert _by_name(index, "past_cost_table.xlsx").sha256 != old_hash
        assert all(Path(d.path).name != "legacy_note.ppt" for d in index.documents())
        # 파일이 사라진 문서는 prune_missing 으로도 제거
        (docs_dir / "template_review.pptx").unlink()
        assert [Path(p).name for p in index.prune_missing()] == ["template_review.pptx"]


def test_only_allowed_root_is_indexed(docs_dir: Path, tmp_path: Path) -> None:
    outside = tmp_path / "밖의 폴더"
    outside.mkdir()
    shutil.copyfile(FIXTURES / "past_review_2023.pptx", outside / "밖.pptx")
    sub = docs_dir / "하위"
    sub.mkdir()
    shutil.copyfile(FIXTURES / "past_review_2023.pptx", sub / "검토_rev7.pptx")
    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        index.index_roots([sub])
        names = {Path(d.path).name for d in index.documents()}
        assert names == {"검토_rev7.pptx"}
        assert _by_name(index, "검토_rev7.pptx").revision == "7"  # 파일명 revision 이 core property 보다 우선
        with pytest.raises(AgentError) as ei:
            index.index_file(outside / "밖.pptx", sub)
        assert ei.value.code == "E_PATH_OUTSIDE_ROOT"
        with pytest.raises(AgentError) as ei:
            index.index_roots([tmp_path / "없는 폴더"])
        assert ei.value.code == "E_PATH_OUTSIDE_ROOT"


@pytest.mark.skipif(not hasattr(os, "symlink") or os.name == "nt", reason="symlink 생성 권한이 필요")
def test_symlinked_dirs_are_not_followed(docs_dir: Path, tmp_path: Path) -> None:
    target = tmp_path / "실제"
    target.mkdir()
    shutil.copyfile(FIXTURES / "past_review_2023.pptx", target / "link_target.pptx")
    try:
        os.symlink(target, docs_dir / "링크", target_is_directory=True)
    except OSError:
        pytest.skip("symlink 생성 불가")
    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        index.index_roots([docs_dir])
        assert all("link_target" not in d.path for d in index.documents())


def test_size_limit_skips_file(docs_dir: Path, tmp_path: Path) -> None:
    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        meta = index.index_file(docs_dir / "past_review_2023.pptx", docs_dir, max_file_mb=0)
        assert meta.status == "SKIPPED_SIZE" and meta.n_fragments == 0


def test_corrupt_file_records_error_not_crash(docs_dir: Path, tmp_path: Path) -> None:
    bad = docs_dir / "손상.xlsx"
    bad.write_bytes(b"not a zip")
    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        rep = index.index_roots([docs_dir])
        assert [Path(e["path"]).name for e in rep.errors] == ["손상.xlsx"]
        assert _by_name(index, "손상.xlsx").status == "ERROR"


def test_scan_package_flags_detects_macro_ole_external_without_executing(tmp_path: Path) -> None:
    p = tmp_path / "flagged.xlsx"
    shutil.copyfile(FIXTURES / "template_comparison.xlsx", p)
    with zipfile.ZipFile(p, "a") as zf:
        zf.writestr("xl/vbaProject.bin", b"\x00fake")
        zf.writestr("xl/externalLinks/externalLink1.xml", "<x/>")
        zf.writestr(
            "xl/worksheets/_rels/sheet9.xml.rels",
            '<Relationships><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="../embeddings/o.bin"/>'
            '<Relationship Type="hyperlink" Target="http://example.invalid" TargetMode="External"/></Relationships>',
        )
    flags = scan_package_flags(p)
    assert {"macro", "ole", "external_link", "hyperlink"} <= set(flags)
    assert scan_package_flags(FIXTURES / "past_review_2023.pptx") == ["chart", "notes"]


def test_helpers() -> None:
    assert revision_from_filename("검토서_rev12.pptx") == "12"
    assert revision_from_filename("cost_v3a.xlsx") == "3a"
    assert revision_from_filename("past_review_2023.pptx") is None
    assert detect_unit("단가 1,234 원 기준") == "원" and detect_unit("중량 12.5 kg") == "kg"
    assert detect_unit("항목") is None
    assert detect_date("검토일 2023-05-01") == "2023-05-01" and detect_date("2023년 5월 1일") == "2023년 5월 1일"
    assert normalize_approval("Approved") == "approved" and normalize_approval("초안") == "draft"
    assert normalize_approval("") == "unknown"


def test_relationship_flags_parse_full_uri_and_package_is_not_ole() -> None:
    from corp_dl_agent.documents.index import relationship_flags

    ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels = (
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{ns}/hyperlink" Target="https://example.invalid" TargetMode="External"/>'
        f'<Relationship Id="rId2" Type="{ns}/oleObject" Target="../embeddings/oleObject1.bin"/>'
        f'<Relationship Id="rId3" TargetMode="External" Type="{ns}/image" Target="file:///C:/x.png"/>'
        f'<Relationship Id="rId4" Type="{ns}/package" Target="../embeddings/Microsoft_Excel_Sheet1.xlsx"/>'
        "</Relationships>"
    )
    assert relationship_flags(rels) == {"hyperlink", "ole", "external_link"}
    # 차트 내장 데이터(package) 만 있는 관계는 flag 없음, 매크로 관계는 macro
    assert relationship_flags(f'<Relationships><Relationship Type="{ns}/package" Target="x.xlsx"/></Relationships>') == set()
    assert relationship_flags(f'<Relationships><Relationship Type="{ns}/vbaProject" Target="vbaProject.bin"/></Relationships>') == {"macro"}


def test_legacy_xls_and_macro_extensions(docs_dir: Path, tmp_path: Path) -> None:
    (docs_dir / "구형 원가표.xls").write_bytes(b"\xd0\xcf\x11\xe0 not parsed")
    shutil.copyfile(FIXTURES / "template_comparison.xlsx", docs_dir / "macro_enabled.xlsm")
    with DocumentIndex(tmp_path / "documents.sqlite") as index:
        rep = index.index_roots([docs_dir])
        assert sorted(Path(p).name for p in rep.needs_conversion) == ["legacy_note.ppt", "구형 원가표.xls"]
        xls = _by_name(index, "구형 원가표.xls")
        assert xls.status == "NEEDS_APPROVED_CONVERSION" and xls.kind == "xls" and xls.n_fragments == 0
        assert "승인" in (xls.error or "")
        macro = _by_name(index, "macro_enabled.xlsm")
        assert macro.status == "INDEXED" and "macro" in macro.flags
        assert rep.flagged[macro.path] == macro.flags
        assert index.fragments(xls.source_id) == []


def test_hidden_shape_attribute_is_flagged_hidden(docs_dir: Path) -> None:
    from pptx import Presentation

    src = docs_dir / "template_review.pptx"
    prs = Presentation(str(src))
    body = next(sh for sh in prs.slides[1].shapes if sh.name == "Body 1")
    for node in body._element.xpath("./*/p:cNvPr"):
        node.set("hidden", "1")
    hidden_path = docs_dir / "숨김 shape.pptx"
    prs.save(str(hidden_path))
    frags, _ = extract_pptx(hidden_path)
    by_loc = {f.locator: f for f in frags}
    assert by_loc["slide:2/shape:Body 1"].hidden is True
    assert by_loc["slide:2/shape:Title 1"].hidden is False
