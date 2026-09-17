"""템플릿 복사본 채우기: 승인된 placeholder / named range / 표 셀만 수정한다.

PPTX (fill_pptx)
- shape 텍스트·표 셀·발표자 노트 안의 ``{{key}}`` placeholder 만 치환한다 (run 서식 보존, run 을 가로지르는 placeholder 는
  같은 paragraph 안에서 병합 후 치환).
- placeholder 문법: ``{{key}}`` 값+단위, ``{{key.value}}`` 숫자만, ``{{key.unit}}``, ``{{key.label}}``, ``{{key.source}}``
  (출처 locator/calculation_id/model_run_id), ``{{key.type}}``. key 는 items, texts, 또는 payload 필드
  (subject, report_id, created_at, data_origin, calculation_version).
- slide 순서·master·theme·차트·그림·OLE 는 그대로 보존한다 (python-pptx 는 패키지를 통째로 round-trip). 차트 데이터는
  payload 로 갱신하지 않으며 unsupported_elements 에 보고한다. 아무것도 삭제하지 않는다.
- 값이 없는 placeholder 는 'MISSING', 상충은 'CONFLICT' 로 표시한다.

XLSX (fill_xlsx)
- 이름이 payload key 와 일치하는 named range(단일 셀) 에만 값을 쓴다. ``key__unit`` / ``key__label`` / ``key__source`` 지원.
  ``tbl_<table>`` 이름의 다중 셀 범위에는 payload.tables[<table>] 의 행을 쓴다 (첫 열 = 행 라벨).
- 수식 셀에는 절대 쓰지 않는다. data_only 로 열지 않으므로 수식은 보존된다.
- openpyxl 은 수식을 계산하지 않고, 저장 시 수식 셀의 cached value 를 버린다 → recalc_status="RECALC_NOT_RUN".
- openpyxl 은 차트/그림/피벗/OLE 를 저장 시 보존하지 못한다. 이런 요소가 있으면 기본적으로 E_DOC_UNSUPPORTED 로 중단하고,
  allow_lossy=True 를 명시했을 때만 호환성 보고(unsupported_elements, preserved=False) 와 함께 진행한다.
"""

from __future__ import annotations

import os
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any, Literal

from pydantic import Field

from corp_dl_agent.common import StrictModel, atomic_write_json, now_iso, sha256_file
from corp_dl_agent.documents.index import iter_pptx_shapes, scan_package_flags, slide_is_hidden
from corp_dl_agent.documents.payload import (
    CONFLICT_TEXT,
    MISSING_TEXT,
    ReportPayload,
    format_number,
    format_quantity,
    provenance_of,
)
from corp_dl_agent.errors import AgentError, blocked_dependency
from corp_dl_agent.version import SCHEMA_VERSION

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)(?:\.(value|unit|label|source|type))?\s*\}\}")
XLSX_NAME_RE = re.compile(r"^([A-Za-z_][A-Za-z0-9_]*?)(?:__(value|unit|label|source|type))?$")
TABLE_RANGE_PREFIX = "tbl_"
PAYLOAD_FIELDS = ("subject", "report_id", "created_at", "data_origin", "calculation_version", "schema_version")
LOSSY_XLSX_FLAGS = ("chart", "drawing", "image", "pivot", "ole")

RecalcStatus = Literal["RECALC_NOT_RUN", "RECALC_DONE"]
RenderStatus = Literal["RENDER_NOT_RUN", "RENDER_DONE"]
ReplacementStatus = Literal["filled", "missing", "conflict", "unknown_key", "unapproved"]


class Replacement(StrictModel):
    locator: str
    placeholder: str
    key: str
    attr: str | None = None
    rendered: str
    status: ReplacementStatus
    numeric: bool = False
    value: float | None = None


class UnsupportedElement(StrictModel):
    locator: str
    element_type: str
    reason_ko: str
    preserved: bool = True


class DocumentManifest(StrictModel):
    schema_version: str = SCHEMA_VERSION
    document_type: Literal["pptx", "xlsx"]
    template_path: str
    input_hash: str
    output_path: str
    output_hash: str
    payload_report_id: str
    created_at: str
    changed_locators: list[str] = Field(default_factory=list)
    replacements: list[Replacement] = Field(default_factory=list)
    unsupported_elements: list[UnsupportedElement] = Field(default_factory=list)
    missing_keys: list[str] = Field(default_factory=list)
    unknown_keys: list[str] = Field(default_factory=list)
    unapproved_keys: list[str] = Field(default_factory=list)
    structure: dict[str, Any] = Field(default_factory=dict)
    flags: list[str] = Field(default_factory=list)
    recalc_status: RecalcStatus = "RECALC_NOT_RUN"
    render_status: RenderStatus = "RENDER_NOT_RUN"
    template_hash_unchanged: bool
    synthetic: bool
    notes: list[str] = Field(default_factory=list)


class GenerationManifest(StrictModel):
    schema_version: str = SCHEMA_VERSION
    created_at: str
    payload_report_id: str
    payload_path: str
    payload_hash: str
    evidence_path: str | None = None
    evidence_hash: str | None = None
    documents: list[DocumentManifest] = Field(default_factory=list)
    recalc_status: RecalcStatus = "RECALC_NOT_RUN"
    render_status: RenderStatus = "RENDER_NOT_RUN"
    synthetic: bool
    notes: list[str] = Field(default_factory=list)


class Resolved(StrictModel):
    key: str
    attr: str | None
    status: ReplacementStatus
    rendered: str
    numeric: bool = False
    value: float | None = None
    text: str | None = None


# ---------------------------------------------------------------------------
# placeholder 해석
# ---------------------------------------------------------------------------


def make_resolver(payload: ReportPayload, approved_keys: set[str] | None = None) -> Callable[[str, str | None], Resolved]:
    """(key, attr) -> Resolved. approved_keys 가 주어지면 그 밖의 key 는 'unapproved' 로 손대지 않는다."""

    def resolve(key: str, attr: str | None) -> Resolved:
        if approved_keys is not None and key not in approved_keys:
            return Resolved(key=key, attr=attr, status="unapproved", rendered="")
        if key in payload.items:
            q = payload.items[key].quantity
            label = payload.items[key].label_ko
            if attr == "unit":
                return Resolved(key=key, attr=attr, status="filled", rendered=q.unit, text=q.unit)
            if attr == "label":
                return Resolved(key=key, attr=attr, status="filled", rendered=label, text=label)
            if attr == "source":
                prov = provenance_of(q)
                return Resolved(key=key, attr=attr, status="filled", rendered=prov, text=prov)
            if attr == "type":
                return Resolved(key=key, attr=attr, status="filled", rendered=q.value_type, text=q.value_type)
            if q.value_type == "conflict":
                return Resolved(key=key, attr=attr, status="conflict", rendered=CONFLICT_TEXT)
            if q.value is None or q.value_type == "missing":
                return Resolved(key=key, attr=attr, status="missing", rendered=MISSING_TEXT)
            if attr == "value":
                return Resolved(
                    key=key, attr=attr, status="filled", rendered=format_number(q.value), numeric=True, value=q.value
                )
            return Resolved(
                key=key,
                attr=attr,
                status="filled",
                rendered=format_quantity(q),
                numeric=True,
                value=q.value,
            )
        if key in payload.texts:
            t = payload.texts[key]
            if attr == "label":
                return Resolved(key=key, attr=attr, status="filled", rendered=t.label_ko, text=t.label_ko)
            if attr == "source":
                src = t.source_locator or t.origin
                return Resolved(key=key, attr=attr, status="filled", rendered=src, text=src)
            if attr in ("unit", "value", "type"):
                return Resolved(key=key, attr=attr, status="filled", rendered=t.text, text=t.text)
            return Resolved(key=key, attr=attr, status="filled", rendered=t.text, text=t.text)
        if key in PAYLOAD_FIELDS and attr is None:
            val = str(getattr(payload, key))
            return Resolved(key=key, attr=attr, status="filled", rendered=val, text=val)
        return Resolved(key=key, attr=attr, status="unknown_key", rendered=MISSING_TEXT)

    return resolve


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def _replace_text(
    text: str, locator: str, resolve: Callable[[str, str | None], Resolved], out: list[Replacement]
) -> str:
    def sub(m: re.Match[str]) -> str:
        key, attr = m.group(1), m.group(2)
        r = resolve(key, attr)
        if r.status == "unapproved":
            out.append(
                Replacement(
                    locator=locator, placeholder=m.group(0), key=key, attr=attr, rendered="", status="unapproved"
                )
            )
            return m.group(0)
        out.append(
            Replacement(
                locator=locator,
                placeholder=m.group(0),
                key=key,
                attr=attr,
                rendered=r.rendered,
                status=r.status,
                numeric=r.numeric,
                value=r.value,
            )
        )
        return r.rendered

    return PLACEHOLDER_RE.sub(sub, text)


def _fill_text_frame(
    tf: Any, locator: str, resolve: Callable[[str, str | None], Resolved], out: list[Replacement]
) -> bool:
    changed = False
    for para in tf.paragraphs:
        if "{{" not in para.text:
            continue
        for run in para.runs:
            if PLACEHOLDER_RE.search(run.text):
                new = _replace_text(run.text, locator, resolve, out)
                if new != run.text:
                    run.text = new
                    changed = True
        if PLACEHOLDER_RE.search(para.text):
            runs = list(para.runs)
            if not runs:
                continue
            merged = "".join(r.text for r in runs)
            if not PLACEHOLDER_RE.search(merged):
                continue
            new = _replace_text(merged, locator, resolve, out)
            if new != merged:
                runs[0].text = new
                for r in runs[1:]:
                    r.text = ""
                changed = True
    return changed


def pptx_structure(prs: Any) -> dict[str, Any]:
    slides: list[dict[str, Any]] = []
    n_charts = n_tables = 0
    for n, slide in enumerate(prs.slides, start=1):
        names: list[str] = []
        for sh, name in iter_pptx_shapes(slide.shapes):
            names.append(name)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                n_charts += 1
            if getattr(sh, "has_table", False) and sh.has_table:
                n_tables += 1
        slides.append(
            {
                "no": n,
                "slide_id": int(slide.slide_id),
                "layout": str(slide.slide_layout.name),
                "hidden": slide_is_hidden(slide),
                "shapes": names,
            }
        )
    return {
        "n_slides": len(slides),
        "slide_ids": [s["slide_id"] for s in slides],
        "n_masters": len(prs.slide_masters),
        "n_layouts": sum(len(m.slide_layouts) for m in prs.slide_masters),
        "n_charts": n_charts,
        "n_tables": n_tables,
        "slides": slides,
    }


def read_pptx_structure(path: str | os.PathLike[str]) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("python-pptx", "documents.templates") from exc
    return pptx_structure(Presentation(str(path)))


def collect_pptx_texts(path: str | os.PathLike[str], *, include_hidden: bool = True) -> dict[str, str]:
    """locator -> 텍스트 (shape/표 셀/노트/차트 계열·분류). 검증·오래된 값 검사용."""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("python-pptx", "documents.templates") from exc
    prs = Presentation(str(path))
    out: dict[str, str] = {}
    for n, slide in enumerate(prs.slides, start=1):
        if slide_is_hidden(slide) and not include_hidden:
            continue
        for sh, name in iter_pptx_shapes(slide.shapes):
            if sh.has_text_frame:
                out[f"slide:{n}/shape:{name}"] = sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for r, row in enumerate(sh.table.rows, start=1):
                    for c, cell in enumerate(row.cells, start=1):
                        out[f"slide:{n}/table:{name}/r{r}c{c}"] = cell.text_frame.text
            if getattr(sh, "has_chart", False) and sh.has_chart:
                chart = sh.chart
                if chart.has_title:
                    out[f"slide:{n}/chart:{name}/title"] = chart.chart_title.text_frame.text
                for plot in chart.plots:
                    cats = [str(c) for c in plot.categories]
                    for series in plot.series:
                        vals = "; ".join(
                            f"{cats[i] if i < len(cats) else i + 1}={v}" for i, v in enumerate(series.values)
                        )
                        out[f"slide:{n}/chart:{name}/series:{series.name}"] = f"{series.name}: {vals}"
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                out[f"slide:{n}/notes"] = tf.text
    return out


def _save_pptx_atomic(prs: Any, out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(f".{out.name}.tmp")
    try:
        prs.save(str(tmp))
        with open(tmp, "rb") as f:
            os.fsync(f.fileno())
        os.replace(tmp, out)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def fill_pptx(
    template: str | os.PathLike[str],
    payload: ReportPayload,
    out: str | os.PathLike[str],
    *,
    approved_keys: set[str] | None = None,
    include_notes: bool = True,
) -> DocumentManifest:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("python-pptx", "documents.templates.fill_pptx") from exc

    template_path = Path(template)
    out_path = Path(out)
    if not template_path.is_file():
        raise AgentError("E_INPUT_INVALID", f"PPTX 템플릿이 없습니다: {template_path.name}")
    if template_path.resolve() == out_path.resolve():
        raise AgentError("E_INPUT_INVALID", "출력 경로가 템플릿과 같습니다. 원본 템플릿은 수정하지 않습니다.")
    input_hash = sha256_file(template_path)
    flags = scan_package_flags(template_path)
    if "macro" in flags:
        raise AgentError(
            "E_DOC_UNSUPPORTED",
            "매크로가 포함된 PPTX 템플릿은 지원하지 않습니다 (매크로는 실행되지 않습니다).",
            details={"template": template_path.name, "flags": flags},
        )
    prs = Presentation(str(template_path))
    resolve = make_resolver(payload, approved_keys)
    replacements: list[Replacement] = []
    changed: list[str] = []
    unsupported: list[UnsupportedElement] = []
    for n, slide in enumerate(prs.slides, start=1):
        for sh, name in iter_pptx_shapes(slide.shapes):
            if sh.has_text_frame:
                loc = f"slide:{n}/shape:{name}"
                if _fill_text_frame(sh.text_frame, loc, resolve, replacements):
                    changed.append(loc)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r, row in enumerate(sh.table.rows, start=1):
                    for c, cell in enumerate(row.cells, start=1):
                        loc = f"slide:{n}/table:{name}/r{r}c{c}"
                        if _fill_text_frame(cell.text_frame, loc, resolve, replacements):
                            changed.append(loc)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                unsupported.append(
                    UnsupportedElement(
                        locator=f"slide:{n}/chart:{name}",
                        element_type="chart",
                        reason_ko="차트 데이터는 payload 로 자동 갱신하지 않습니다 (원본 그대로 보존, 수동 확인 필요)",
                    )
                )
            elif sh.shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                unsupported.append(
                    UnsupportedElement(
                        locator=f"slide:{n}/shape:{name}",
                        element_type="ole",
                        reason_ko="OLE 개체는 실행/갱신하지 않고 보존합니다",
                    )
                )
            elif sh.shape_type in (MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.WEB_VIDEO):
                unsupported.append(
                    UnsupportedElement(
                        locator=f"slide:{n}/shape:{name}",
                        element_type="media",
                        reason_ko="미디어 개체는 갱신하지 않고 보존합니다",
                    )
                )
            elif (
                not sh.has_text_frame
                and not (getattr(sh, "has_table", False) and sh.has_table)
                and sh._element.tag.endswith("}graphicFrame")
            ):
                unsupported.append(
                    UnsupportedElement(
                        locator=f"slide:{n}/shape:{name}",
                        element_type="graphic_frame",
                        reason_ko="SmartArt 등 graphicFrame 은 갱신하지 않고 보존합니다",
                    )
                )
        if include_notes and slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                loc = f"slide:{n}/notes"
                if _fill_text_frame(tf, loc, resolve, replacements):
                    changed.append(loc)
    structure = pptx_structure(prs)
    _save_pptx_atomic(prs, out_path)
    notes = [
        "slide 순서·master·theme·차트·그림은 원본 그대로 보존됩니다 (python-pptx 패키지 round-trip).",
        "렌더링(잘림/겹침/폰트) 검사는 수행하지 않았습니다: RENDER_NOT_RUN.",
    ]
    return _manifest(
        "pptx", template_path, input_hash, out_path, payload, changed, replacements, unsupported, structure, flags, notes
    )


def _manifest(
    doc_type: Literal["pptx", "xlsx"],
    template_path: Path,
    input_hash: str,
    out_path: Path,
    payload: ReportPayload,
    changed: list[str],
    replacements: list[Replacement],
    unsupported: list[UnsupportedElement],
    structure: dict[str, Any],
    flags: list[str],
    notes: list[str],
) -> DocumentManifest:
    missing = sorted({r.key for r in replacements if r.status in ("missing", "conflict")})
    unknown = sorted({r.key for r in replacements if r.status == "unknown_key"})
    unapproved = sorted({r.key for r in replacements if r.status == "unapproved"})
    if "external_link" in flags:
        notes.append("외부 링크가 있는 템플릿입니다. 링크는 갱신/실행하지 않았습니다.")
    return DocumentManifest(
        document_type=doc_type,
        template_path=str(template_path),
        input_hash=input_hash,
        output_path=str(out_path),
        output_hash=sha256_file(out_path),
        payload_report_id=payload.report_id,
        created_at=now_iso(),
        changed_locators=sorted(set(changed), key=changed.index),
        replacements=replacements,
        unsupported_elements=unsupported,
        missing_keys=missing,
        unknown_keys=unknown,
        unapproved_keys=unapproved,
        structure=structure,
        flags=flags,
        template_hash_unchanged=sha256_file(template_path) == input_hash,
        synthetic=payload.synthetic,
        notes=notes,
    )


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def xlsx_structure(wb: Any) -> dict[str, Any]:
    formulas: dict[str, str] = {}
    tables: dict[str, dict[str, str]] = {}
    merged: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas[f"{ws.title}!{cell.coordinate}"] = str(cell.value)
        tables[ws.title] = {str(k): str(v) for k, v in ws.tables.items()}
        merged[ws.title] = sorted(str(r) for r in ws.merged_cells.ranges)
    names = sorted(str(k) for k in wb.defined_names.keys())
    for ws in wb.worksheets:
        names.extend(sorted(f"{ws.title}!{k}" for k in ws.defined_names.keys()))
    return {
        "sheets": list(wb.sheetnames),
        "sheet_states": {ws.title: str(ws.sheet_state) for ws in wb.worksheets},
        "defined_names": names,
        "formulas": formulas,
        "tables": tables,
        "merged": merged,
    }


def read_xlsx_structure(path: str | os.PathLike[str]) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("openpyxl", "documents.templates") from exc
    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        return xlsx_structure(wb)
    finally:
        wb.close()


def collect_xlsx_texts(path: str | os.PathLike[str], *, include_hidden: bool = True) -> dict[str, str]:
    """locator -> 셀 텍스트(수식은 원문). named range 단일 셀은 range:<name> 으로도 반환."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("openpyxl", "documents.templates") from exc
    wb = openpyxl.load_workbook(str(path), data_only=False)
    out: dict[str, str] = {}
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible" and not include_hidden:
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    out[f"sheet:{ws.title}!{cell.coordinate}"] = _xl_text(cell.value)
        for name, dn in _all_defined_names(wb):
            for sheet, coord in _destinations(dn):
                c = coord.replace("$", "")
                if ":" in c or sheet not in wb.sheetnames:
                    continue
                v = wb[sheet][c].value
                out[f"range:{name}"] = _xl_text(v) if v is not None else ""
    finally:
        wb.close()
    return out


def _xl_text(v: Any) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _all_defined_names(wb: Any) -> list[tuple[str, Any]]:
    names: list[tuple[str, Any]] = [(str(k), v) for k, v in wb.defined_names.items()]
    for ws in wb.worksheets:
        names.extend((str(k), v) for k, v in ws.defined_names.items())
    return names


def _destinations(dn: Any) -> list[tuple[str, str]]:
    try:
        return [(str(s), str(c)) for s, c in dn.destinations]
    except Exception:  # noqa: BLE001 - 상수/외부 이름
        return []


def _is_formula_cell(cell: Any) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def fill_xlsx(
    template: str | os.PathLike[str],
    payload: ReportPayload,
    out: str | os.PathLike[str],
    *,
    approved_names: set[str] | None = None,
    allow_lossy: bool = False,
) -> DocumentManifest:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("openpyxl", "documents.templates.fill_xlsx") from exc

    template_path = Path(template)
    out_path = Path(out)
    if not template_path.is_file():
        raise AgentError("E_INPUT_INVALID", f"XLSX 템플릿이 없습니다: {template_path.name}")
    if template_path.resolve() == out_path.resolve():
        raise AgentError("E_INPUT_INVALID", "출력 경로가 템플릿과 같습니다. 원본 템플릿은 수정하지 않습니다.")
    input_hash = sha256_file(template_path)
    flags = scan_package_flags(template_path)
    if "macro" in flags:
        raise AgentError(
            "E_DOC_UNSUPPORTED",
            "매크로가 포함된 XLSX 템플릿은 지원하지 않습니다 (매크로는 실행되지 않습니다).",
            details={"template": template_path.name, "flags": flags},
        )
    lossy = [f for f in flags if f in LOSSY_XLSX_FLAGS]
    unsupported: list[UnsupportedElement] = []
    if lossy:
        if not allow_lossy:
            raise AgentError(
                "E_DOC_UNSUPPORTED",
                f"XLSX 템플릿에 openpyxl 이 보존하지 못하는 요소가 있습니다: {', '.join(lossy)}",
                hint="차트/그림/피벗/OLE 를 별도 시트·파일로 분리하거나, 호환성 보고를 감수하고 --allow-lossy-xlsx 로 진행하세요. "
                "대체 경로: 승인된 Excel 엔진(excel_com)에서 값만 입력.",
                details={"template": template_path.name, "lossy_elements": lossy},
            )
        for f in lossy:
            unsupported.append(
                UnsupportedElement(
                    locator="workbook",
                    element_type=f,
                    reason_ko="openpyxl 저장 시 보존되지 않는 요소 (LOST_ON_SAVE). allow_lossy 로 명시 진행됨",
                    preserved=False,
                )
            )
    wb = openpyxl.load_workbook(str(template_path), data_only=False)
    try:
        before = xlsx_structure(wb)
        resolve = make_resolver(payload, approved_names)
        replacements: list[Replacement] = []
        changed: list[str] = []
        notes: list[str] = []
        for name, dn in _all_defined_names(wb):
            dests = _destinations(dn)
            if not dests:
                continue
            for sheet, coord in dests:
                if sheet not in wb.sheetnames:
                    continue
                ws = wb[sheet]
                c = coord.replace("$", "")
                loc = f"range:{name}"
                if ":" in c:
                    if name.startswith(TABLE_RANGE_PREFIX):
                        _fill_table_range(ws, c, name[len(TABLE_RANGE_PREFIX) :], payload, loc, replacements, changed, unsupported)
                    continue
                cell = ws[c]
                if _is_formula_cell(cell):
                    m = XLSX_NAME_RE.match(name)
                    if m and (m.group(1) in payload.items or m.group(1) in payload.texts):
                        unsupported.append(
                            UnsupportedElement(
                                locator=loc,
                                element_type="formula_cell",
                                reason_ko="named range 가 수식 셀을 가리켜 값을 쓰지 않았습니다 (수식 보존)",
                            )
                        )
                    continue
                m = XLSX_NAME_RE.match(name)
                if not m:
                    continue
                key, attr = m.group(1), m.group(2)
                r = resolve(key, attr)
                if r.status == "unapproved":
                    replacements.append(
                        Replacement(locator=loc, placeholder=name, key=key, attr=attr, rendered="", status="unapproved")
                    )
                    continue
                if r.status == "unknown_key":
                    if cell.value is None:
                        cell.value = MISSING_TEXT
                        changed.append(loc)
                        replacements.append(
                            Replacement(
                                locator=loc, placeholder=name, key=key, attr=attr, rendered=MISSING_TEXT, status="unknown_key"
                            )
                        )
                    continue
                if r.numeric and r.value is not None:
                    cell.value = int(r.value) if float(r.value).is_integer() else float(r.value)
                    rendered = format_number(r.value)
                else:
                    cell.value = r.rendered if r.text is None else r.text
                    rendered = r.rendered
                changed.append(loc)
                replacements.append(
                    Replacement(
                        locator=loc,
                        placeholder=name,
                        key=key,
                        attr=attr,
                        rendered=rendered,
                        status=r.status,
                        numeric=r.numeric,
                        value=r.value,
                    )
                )
        after = xlsx_structure(wb)
        if after["formulas"] != before["formulas"]:
            raise AgentError(
                "E_DOC_STRUCTURE_CHANGED",
                "수식 셀이 변경되었습니다. named range 가 수식 셀을 가리키는지 확인하세요.",
                details={"changed": sorted(set(before["formulas"]) ^ set(after["formulas"]))[:20]},
            )
        out_path.parent.mkdir(parents=True, exist_ok=True)
        tmp = out_path.with_name(f".{out_path.name}.tmp")
        try:
            wb.save(str(tmp))
            with open(tmp, "rb") as f:
                os.fsync(f.fileno())
            os.replace(tmp, out_path)
        except BaseException:
            try:
                os.unlink(tmp)
            except OSError:
                pass
            raise
    finally:
        wb.close()
    notes.extend(
        [
            "openpyxl 은 수식을 계산하지 않으며 저장 시 수식 셀의 cached value 를 제거합니다: RECALC_NOT_RUN "
            "(Excel 또는 승인된 재계산 엔진에서 열어 재계산 필요).",
            "인쇄영역/렌더링 검사는 수행하지 않았습니다: RENDER_NOT_RUN.",
            f"수식 셀 {len(after['formulas'])}개 보존.",
        ]
    )
    return _manifest(
        "xlsx", template_path, input_hash, out_path, payload, changed, replacements, unsupported, after, flags, notes
    )


def _fill_table_range(
    ws: Any,
    ref: str,
    table_name: str,
    payload: ReportPayload,
    loc: str,
    replacements: list[Replacement],
    changed: list[str],
    unsupported: list[UnsupportedElement],
) -> None:
    from openpyxl.utils.cell import range_boundaries

    table = payload.tables.get(table_name)
    if table is None:
        return
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    c0, r0, c1, r1 = int(min_col or 1), int(min_row or 1), int(max_col or 1), int(max_row or 1)
    n_rows_avail = r1 - r0 + 1
    n_cols_avail = c1 - c0 + 1
    if len(table.rows) > n_rows_avail or len(table.columns) + 1 > n_cols_avail:
        unsupported.append(
            UnsupportedElement(
                locator=loc,
                element_type="table_range_too_small",
                reason_ko=f"payload 표 '{table_name}' ({len(table.rows)}행 x {len(table.columns) + 1}열) 가 범위 {ref} 보다 큽니다. 맞는 부분만 기록했습니다",
            )
        )
    for i, row in enumerate(table.rows[:n_rows_avail]):
        r = r0 + i
        label_cell = ws.cell(row=r, column=c0)
        if not _is_formula_cell(label_cell):
            label_cell.value = row.label_ko
        for j, col in enumerate(table.columns[: n_cols_avail - 1]):
            cell = ws.cell(row=r, column=c0 + 1 + j)
            cell_loc = f"{loc}/r{i + 1}c{j + 2}"
            if _is_formula_cell(cell):
                unsupported.append(
                    UnsupportedElement(
                        locator=cell_loc, element_type="formula_cell", reason_ko="표 범위 안의 수식 셀은 쓰지 않습니다"
                    )
                )
                continue
            q = row.cells.get(col.key)
            if q is None or q.value_type == "conflict":
                cell.value = CONFLICT_TEXT if q is not None else MISSING_TEXT
                status: ReplacementStatus = "conflict" if q is not None else "missing"
                replacements.append(
                    Replacement(
                        locator=cell_loc, placeholder=f"{table_name}.{col.key}", key=table_name, attr=col.key, rendered=str(cell.value), status=status
                    )
                )
            elif q.value is None or q.value_type == "missing":
                cell.value = MISSING_TEXT
                replacements.append(
                    Replacement(
                        locator=cell_loc, placeholder=f"{table_name}.{col.key}", key=table_name, attr=col.key, rendered=MISSING_TEXT, status="missing"
                    )
                )
            else:
                cell.value = int(q.value) if float(q.value).is_integer() else float(q.value)
                replacements.append(
                    Replacement(
                        locator=cell_loc,
                        placeholder=f"{table_name}.{col.key}",
                        key=table_name,
                        attr=col.key,
                        rendered=format_number(q.value),
                        status="filled",
                        numeric=True,
                        value=q.value,
                    )
                )
            changed.append(cell_loc)


# ---------------------------------------------------------------------------
# manifest 파일
# ---------------------------------------------------------------------------


def write_generation_manifest(manifest: GenerationManifest, path: str | os.PathLike[str]) -> Path:
    p = Path(path)
    atomic_write_json(p, manifest.model_dump(mode="json"))
    return p
