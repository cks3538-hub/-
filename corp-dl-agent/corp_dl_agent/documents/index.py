"""문서 색인: 허용된 sources root 의 PPTX/XLSX 를 sqlite 에 색인한다.

locator 규칙
- PPTX: ``slide:3/shape:Title 1`` (shape 텍스트, group 안이면 ``Group 2/TextBox 1``),
        ``slide:3/table:Table 1/r2c1`` (표 셀, 1-based, r1 = 머리글),
        ``slide:3/notes`` (발표자 노트), ``slide:3/chart:Chart 1/series:중량`` (차트 계열 원데이터)
- XLSX: ``sheet:Cost!B7`` (셀), ``range:UnitCost`` (named range), ``table:T1!r2c3`` (ListObject 셀, r1 = 머리글)

보존: source_id(경로 기반 안정 ID), sha256, revision(파일명 > core properties), mtime, hidden 여부,
수식(formula) 과 cached value 의 구분, 단위/기준일/approval_state 추정.
구형 .ppt/.xls 는 파싱하지 않고 ``NEEDS_APPROVED_CONVERSION`` 상태로만 기록한다.
매크로/OLE/외부 링크는 실행하지 않고 zip 항목·관계 파일을 읽어 flag 만 남긴다.
삭제된 파일은 색인에서 제거한다(prune).
"""

from __future__ import annotations

import json
import os
import re
import sqlite3
import zipfile
from collections.abc import Iterable, Iterator
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

from pydantic import Field

from corp_dl_agent.common import StrictModel, now_iso, sha256_file, stable_id
from corp_dl_agent.errors import AgentError, blocked_dependency
from corp_dl_agent.security.paths import check_zip_safety, has_link_component

INDEX_SCHEMA_VERSION = 1

# 확장자 -> 문서 종류. 매크로 사용 파일(.pptm/.xlsm) 은 색인은 시도하되 flag 를 남긴다.
SUPPORTED_EXT: dict[str, str] = {".pptx": "pptx", ".pptm": "pptx", ".xlsx": "xlsx", ".xlsm": "xlsx"}
LEGACY_EXT: dict[str, str] = {".ppt": "ppt", ".xls": "xls"}

DocStatus = Literal["INDEXED", "UNCHANGED", "NEEDS_APPROVED_CONVERSION", "SKIPPED_SIZE", "ERROR"]
DocKind = Literal["form_reference", "fact"]
FragmentKind = Literal[
    "slide_text", "table_cell", "notes", "chart_series", "chart_title", "cell", "named_range", "xl_table_cell"
]

PLACEHOLDER_MARK = "{{"
FORM_NAME_MARKERS = ("template", "양식", "서식", "form_", "_form", "templ")
FORM_DIR_MARKERS = ("templates", "template", "양식", "서식")

_UNIT_RE = re.compile(
    r"(?<![\w.])-?\d[\d,]*(?:\.\d+)?\s*(kg|g|mg|t|mm|cm|m|N|kN|MPa|원|KRW|USD|%|ea|EA|개|초|s|h)(?![\w])"
)
_DATE_RES = (
    re.compile(r"(?<!\d)(\d{4}-\d{2}-\d{2})(?!\d)"),
    re.compile(r"(?<!\d)(\d{4}\.\d{2}\.\d{2})(?!\d)"),
    re.compile(r"(\d{4}년\s*\d{1,2}월\s*\d{1,2}일)"),
)
_REV_RE = re.compile(r"(?i)(?:^|[_\-\s(\[])(?:rev|r|v)[\-_ ]?(\d+[a-z]?)(?=$|[_\-\s)\].])")
_NUMBER_FORMAT_UNIT_RE = re.compile(r'"([^"]+)"')
# OOXML 관계(.rels) 파일: 실행하지 않고 <Relationship .../> 의 Type/TargetMode 속성만 읽는다
_REL_TAG_RE = re.compile(r"<Relationship\b([^>]*)/?>", re.IGNORECASE)
_REL_ATTR_RE = re.compile(r'([A-Za-z:]+)\s*=\s*"([^"]*)"')
# 관계 Type 의 마지막 경로 구분(소문자) -> flag. 전체 URI 이든 축약형이든 마지막 구분만 본다.
_REL_TYPE_FLAGS: dict[str, str] = {
    "oleobject": "ole",  # ('package' 관계는 차트의 내장 데이터 통합문서 등 정상 요소이므로 flag 하지 않는다)
    "vbaproject": "macro",
    "hyperlink": "hyperlink",
    "externallink": "external_link",
    "externallinkpath": "external_link",
}


def relationship_flags(rels_xml: str) -> set[str]:
    """관계 XML 문자열에서 macro/ole/hyperlink/external_link flag 를 추출한다 (실행 없음)."""
    flags: set[str] = set()
    for m in _REL_TAG_RE.finditer(rels_xml):
        attrs = {k.lower(): v for k, v in _REL_ATTR_RE.findall(m.group(1))}
        rel_type = attrs.get("type", "").rstrip("/").rsplit("/", 1)[-1].lower()
        flag = _REL_TYPE_FLAGS.get(rel_type)
        if flag:
            flags.add(flag)
        if attrs.get("targetmode", "").lower() == "external" and flag != "hyperlink":
            flags.add("external_link")
    return flags


class Fragment(StrictModel):
    """색인 단위(검색·근거 인용 단위)."""

    locator: str
    fragment_kind: FragmentKind
    text: str
    hidden: bool = False
    formula: str | None = None
    cached_value: str | None = None
    number_format: str | None = None
    unit: str | None = None
    base_date: str | None = None
    meta: dict[str, str] = Field(default_factory=dict)


class DocumentMeta(StrictModel):
    source_id: str
    path: str
    root: str
    kind: str
    status: DocStatus
    sha256: str
    size: int
    mtime: float
    revision: str
    title: str
    author: str
    created: str | None = None
    modified: str | None = None
    approval_state: str = "unknown"
    doc_kind: DocKind = "fact"
    synthetic: bool = False
    flags: list[str] = Field(default_factory=list)
    n_fragments: int = 0
    hidden_fragments: int = 0
    indexed_at: str = ""
    error: str | None = None


class IndexReport(StrictModel):
    roots: list[str]
    started_at: str
    finished_at: str
    indexed: list[str] = Field(default_factory=list)
    unchanged: list[str] = Field(default_factory=list)
    needs_conversion: list[str] = Field(default_factory=list)
    skipped: list[dict[str, str]] = Field(default_factory=list)
    errors: list[dict[str, str]] = Field(default_factory=list)
    pruned: list[str] = Field(default_factory=list)
    flagged: dict[str, list[str]] = Field(default_factory=dict)
    n_documents: int = 0
    n_fragments: int = 0
    n_hidden_fragments: int = 0


# ---------------------------------------------------------------------------
# 텍스트 보조
# ---------------------------------------------------------------------------


def detect_unit(text: str) -> str | None:
    m = _UNIT_RE.search(text)
    return m.group(1) if m else None


def detect_date(text: str) -> str | None:
    for pat in _DATE_RES:
        m = pat.search(text)
        if m:
            return m.group(1)
    return None


def revision_from_filename(name: str) -> str | None:
    stem = Path(name).stem
    m = _REV_RE.search(stem)
    return m.group(1) if m else None


def normalize_approval(raw: str | None) -> str:
    v = (raw or "").strip().lower()
    if not v:
        return "unknown"
    if v in ("approved", "final", "released") or "승인" in v or "확정" in v:
        return "approved"
    if v in ("draft", "wip", "in review", "review") or "초안" in v or "검토" in v:
        return "draft"
    if v in ("obsolete", "superseded", "retired") or "폐기" in v or "구버전" in v:
        return "obsolete"
    return v


def _is_synthetic_props(*fields: str | None) -> bool:
    return any("synthetic" in (f or "").lower() or "합성" in (f or "") for f in fields)


def _iso(dt: Any) -> str | None:
    if isinstance(dt, datetime):
        return dt.isoformat(timespec="seconds")
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.isoformat(timespec="seconds")
            if (value.hour or value.minute or value.second)
            else value.date().isoformat()
        )
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# ---------------------------------------------------------------------------
# 패키지(zip) flag 검사: 실행하지 않고 항목 이름/관계 파일만 읽는다
# ---------------------------------------------------------------------------


def scan_package_flags(path: Path, *, max_members: int = 20000) -> list[str]:
    """OOXML zip 의 매크로/OLE/외부 링크/차트/그림/피벗 존재 여부. 어떤 콘텐츠도 실행하지 않는다."""
    flags: set[str] = set()
    with zipfile.ZipFile(path) as zf:
        names = check_zip_safety(zf, max_members=max_members)
        for n in names:
            low = n.lower()
            if "vbaproject" in low:
                flags.add("macro")
            if low.startswith("xl/externallinks/"):
                flags.add("external_link")
            if low.startswith(("ppt/charts/", "xl/charts/")):
                flags.add("chart")
            if low.startswith(("xl/drawings/", "ppt/drawings/")):
                flags.add("drawing")
            if low.startswith(("xl/media/", "ppt/media/")):
                flags.add("image")
            if low.startswith("xl/pivot"):
                flags.add("pivot")
            if low.startswith(("xl/comments", "ppt/comments/")):
                flags.add("comments")
            if low.startswith("ppt/notesslides/"):
                flags.add("notes")
            if low.endswith(".rels"):
                try:
                    rel = zf.read(n).decode("utf-8", errors="replace")
                except (KeyError, OSError, zipfile.BadZipFile):
                    continue
                flags |= relationship_flags(rel)
    if path.suffix.lower() in (".pptm", ".xlsm"):
        flags.add("macro")
    return sorted(flags)


# ---------------------------------------------------------------------------
# PPTX 추출
# ---------------------------------------------------------------------------


def _shape_hidden(shape: Any) -> bool:
    try:
        nodes = shape._element.xpath("./*/p:cNvPr")
    except Exception:  # noqa: BLE001 - lxml/xpath 실패 시 hidden 아님으로 간주
        return False
    for node in nodes:
        if str(node.get("hidden", "0")).lower() in ("1", "true"):
            return True
    return False


def iter_pptx_shapes(shapes: Any, prefix: str = "") -> Iterator[tuple[Any, str]]:
    """group 을 재귀적으로 펼치며 (shape, 이름 경로) 를 돌려준다."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for sh in shapes:
        name = f"{prefix}{sh.name}"
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_pptx_shapes(sh.shapes, prefix=f"{name}/")
        else:
            yield sh, name


def slide_is_hidden(slide: Any) -> bool:
    return str(slide._element.get("show", "1")) == "0"


def _text_fragment(locator: str, kind: FragmentKind, text: str, hidden: bool, **meta: str) -> Fragment:
    return Fragment(
        locator=locator,
        fragment_kind=kind,
        text=text,
        hidden=hidden,
        unit=detect_unit(text),
        base_date=detect_date(text),
        meta={k: v for k, v in meta.items() if v},
    )


def extract_pptx(path: Path, *, include_notes: bool = True) -> tuple[list[Fragment], dict[str, Any]]:
    """PPTX 의 slide/shape/table/notes/chart 원데이터를 Fragment 목록으로. 두 번째 반환값은 core properties."""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - 설치 환경에서는 도달하지 않음
        raise blocked_dependency("python-pptx", "documents.index") from exc

    prs = Presentation(str(path))
    cp = prs.core_properties
    props: dict[str, Any] = {
        "title": cp.title or "",
        "author": cp.author or "",
        "revision": str(cp.revision) if cp.revision else "",
        "created": _iso(cp.created),
        "modified": _iso(cp.modified),
        "category": cp.category or "",
        "keywords": cp.keywords or "",
        "content_status": cp.content_status or "",
        "comments": cp.comments or "",
        "n_slides": len(prs.slides),
        "hidden_slides": 0,
    }
    frags: list[Fragment] = []
    for n, slide in enumerate(prs.slides, start=1):
        s_hidden = slide_is_hidden(slide)
        if s_hidden:
            props["hidden_slides"] += 1
        for sh, name in iter_pptx_shapes(slide.shapes):
            hidden = s_hidden or _shape_hidden(sh)
            if sh.has_text_frame:
                text = sh.text_frame.text.strip()
                if text:
                    frags.append(_text_fragment(f"slide:{n}/shape:{name}", "slide_text", text, hidden))
            if getattr(sh, "has_table", False) and sh.has_table:
                for r, row in enumerate(sh.table.rows, start=1):
                    for c, cell in enumerate(row.cells, start=1):
                        text = cell.text_frame.text.strip()
                        if text:
                            frags.append(
                                _text_fragment(f"slide:{n}/table:{name}/r{r}c{c}", "table_cell", text, hidden)
                            )
            if getattr(sh, "has_chart", False) and sh.has_chart:
                frags.extend(_chart_fragments(sh.chart, n, name, hidden))
        if include_notes and slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
            if text:
                frags.append(_text_fragment(f"slide:{n}/notes", "notes", text, s_hidden))
    return frags, props


def _chart_fragments(chart: Any, n: int, name: str, hidden: bool) -> list[Fragment]:
    out: list[Fragment] = []
    chart_type = str(getattr(chart, "chart_type", "") or "")
    if chart.has_title:
        title = chart.chart_title.text_frame.text.strip()
        if title:
            out.append(
                _text_fragment(
                    f"slide:{n}/chart:{name}/title", "chart_title", title, hidden, chart_type=chart_type
                )
            )
    for plot in chart.plots:
        cats = [str(c) for c in plot.categories]
        for series in plot.series:
            values = list(series.values)
            pairs = []
            for i, v in enumerate(values):
                cat = cats[i] if i < len(cats) else f"#{i + 1}"
                pairs.append(f"{cat}={_cell_text(v)}")
            sname = str(series.name or f"series{series.index + 1}")
            text = f"{sname}: " + "; ".join(pairs)
            frag = _text_fragment(
                f"slide:{n}/chart:{name}/series:{sname}",
                "chart_series",
                text,
                hidden,
                chart_type=chart_type,
                categories=json.dumps(cats, ensure_ascii=False),
                values=json.dumps([None if v is None else float(v) for v in values]),
            )
            out.append(frag)
    return out


# ---------------------------------------------------------------------------
# XLSX 추출
# ---------------------------------------------------------------------------


def _range_cells(ref: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
    return int(min_col or 1), int(min_row or 1), int(max_col or 1), int(max_row or 1)


def extract_xlsx(path: Path) -> tuple[list[Fragment], dict[str, Any]]:
    """XLSX 의 sheet/cell/named range/table 을 Fragment 목록으로. 수식과 cached value 를 구분한다.

    openpyxl 은 수식을 계산하지 않는다: formula 는 원문, cached_value 는 파일에 저장된 마지막 계산값(없을 수 있음).
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("openpyxl", "documents.index") from exc

    wb_f = openpyxl.load_workbook(str(path), data_only=False, keep_links=False)
    wb_v = openpyxl.load_workbook(str(path), data_only=True, keep_links=False)
    p = wb_f.properties
    props: dict[str, Any] = {
        "title": p.title or "",
        "author": p.creator or "",
        "revision": str(p.revision) if p.revision else "",
        "created": _iso(p.created),
        "modified": _iso(p.modified),
        "category": p.category or "",
        "keywords": p.keywords or "",
        "content_status": p.contentStatus or "",
        "comments": p.description or "",
        "sheets": list(wb_f.sheetnames),
        "hidden_sheets": [ws.title for ws in wb_f.worksheets if ws.sheet_state != "visible"],
    }
    frags: list[Fragment] = []
    cell_lookup: dict[tuple[str, str], Fragment] = {}
    for ws in wb_f.worksheets:
        ws_v = wb_v[ws.title]
        sheet_hidden = ws.sheet_state != "visible"
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                coord = str(cell.coordinate)
                col_letter = str(cell.column_letter)
                hidden = (
                    sheet_hidden
                    or bool(ws.row_dimensions[cell.row].hidden)
                    or bool(ws.column_dimensions[col_letter].hidden)
                )
                frag = _xlsx_cell_fragment(
                    f"sheet:{ws.title}!{coord}",
                    "cell",
                    cell,
                    ws_v[coord].value,
                    hidden,
                    sheet=ws.title,
                    cell=coord,
                )
                frags.append(frag)
                cell_lookup[(ws.title, coord)] = frag
        for tname, ref in ws.tables.items():
            c0, r0, c1, r1 = _range_cells(str(ref))
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    cell = ws.cell(row=r, column=c)
                    if cell.value is None:
                        continue
                    coord = str(cell.coordinate)
                    base = cell_lookup.get((ws.title, coord))
                    hidden = base.hidden if base else sheet_hidden
                    frags.append(
                        _xlsx_cell_fragment(
                            f"table:{tname}!r{r - r0 + 1}c{c - c0 + 1}",
                            "xl_table_cell",
                            cell,
                            ws_v[coord].value,
                            hidden,
                            sheet=ws.title,
                            cell=coord,
                            table=str(tname),
                        )
                    )
    # named ranges (workbook 범위 + sheet 범위)
    named: list[tuple[str, Any, str | None]] = [(str(k), v, None) for k, v in wb_f.defined_names.items()]
    for ws in wb_f.worksheets:
        named.extend((str(k), v, ws.title) for k, v in ws.defined_names.items())
    for name, dn, scope in named:
        try:
            dests = list(dn.destinations)
        except Exception:  # noqa: BLE001 - 상수/외부 참조 이름은 destinations 가 없을 수 있음
            dests = []
        dn_hidden = bool(getattr(dn, "hidden", False))
        if not dests:
            text = str(getattr(dn, "attr_text", "") or "")
            frags.append(
                Fragment(
                    locator=f"range:{name}",
                    fragment_kind="named_range",
                    text=text,
                    hidden=dn_hidden,
                    meta={"scope": scope or "workbook", "kind": "constant_or_external"},
                )
            )
            continue
        for sheet_title, coord in dests:
            coord_clean = str(coord).replace("$", "")
            if sheet_title not in wb_f.sheetnames:
                continue
            if ":" not in coord_clean:
                base = cell_lookup.get((sheet_title, coord_clean))
                ws = wb_f[sheet_title]
                cell = ws[coord_clean]
                frag = _xlsx_cell_fragment(
                    f"range:{name}",
                    "named_range",
                    cell,
                    wb_v[sheet_title][coord_clean].value,
                    (base.hidden if base else ws.sheet_state != "visible") or dn_hidden,
                    sheet=sheet_title,
                    cell=coord_clean,
                    scope=scope or "workbook",
                )
                frags.append(frag)
            else:
                c0, r0, c1, r1 = _range_cells(coord_clean)
                frags.append(
                    Fragment(
                        locator=f"range:{name}",
                        fragment_kind="named_range",
                        text=f"{sheet_title}!{coord_clean} ({(r1 - r0 + 1) * (c1 - c0 + 1)} cells)",
                        hidden=dn_hidden or wb_f[sheet_title].sheet_state != "visible",
                        meta={"sheet": sheet_title, "ref": coord_clean, "scope": scope or "workbook"},
                    )
                )
    wb_f.close()
    wb_v.close()
    return frags, props


def _xlsx_cell_fragment(
    locator: str, kind: FragmentKind, xl_cell: Any, cached: Any, hidden: bool, **meta: str
) -> Fragment:
    value = xl_cell.value
    number_format = (
        str(xl_cell.number_format) if xl_cell.number_format and xl_cell.number_format != "General" else None
    )
    formula: str | None = None
    cached_value: str | None = None
    if xl_cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        formula = str(value)
        cached_value = _cell_text(cached) if cached is not None else None
        text = formula if cached_value is None else f"{formula} = {cached_value}"
    else:
        text = _cell_text(value)
    unit = detect_unit(text)
    if unit is None and number_format:
        m = _NUMBER_FORMAT_UNIT_RE.search(number_format)
        if m:
            unit = m.group(1).strip() or None
    return Fragment(
        locator=locator,
        fragment_kind=kind,
        text=text,
        hidden=hidden,
        formula=formula,
        cached_value=cached_value,
        number_format=number_format,
        unit=unit,
        base_date=detect_date(text),
        meta={k: v for k, v in meta.items() if v},
    )


# ---------------------------------------------------------------------------
# 문서 분류
# ---------------------------------------------------------------------------


def classify_doc_kind(path: Path, root: Path, fragments: Iterable[Fragment]) -> DocKind:
    """과거 양식(form_reference) vs 현재 사실 근거(fact). placeholder 가 있거나 이름/폴더가 양식이면 form_reference."""
    low = path.name.lower()
    if any(m in low for m in FORM_NAME_MARKERS):
        return "form_reference"
    try:
        rel_parts = path.relative_to(root).parts[:-1]
    except ValueError:
        rel_parts = ()
    if any(part.lower() in FORM_DIR_MARKERS for part in rel_parts):
        return "form_reference"
    if any(PLACEHOLDER_MARK in f.text for f in fragments):
        return "form_reference"
    return "fact"


# ---------------------------------------------------------------------------
# sqlite 색인
# ---------------------------------------------------------------------------

_SCHEMA = """
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS documents (
  source_id TEXT PRIMARY KEY,
  path TEXT NOT NULL UNIQUE,
  root TEXT NOT NULL,
  kind TEXT NOT NULL,
  status TEXT NOT NULL,
  sha256 TEXT NOT NULL,
  size INTEGER NOT NULL,
  mtime REAL NOT NULL,
  revision TEXT NOT NULL DEFAULT '',
  title TEXT NOT NULL DEFAULT '',
  author TEXT NOT NULL DEFAULT '',
  created TEXT,
  modified TEXT,
  approval_state TEXT NOT NULL DEFAULT 'unknown',
  doc_kind TEXT NOT NULL DEFAULT 'fact',
  synthetic INTEGER NOT NULL DEFAULT 0,
  flags TEXT NOT NULL DEFAULT '[]',
  n_fragments INTEGER NOT NULL DEFAULT 0,
  hidden_fragments INTEGER NOT NULL DEFAULT 0,
  indexed_at TEXT NOT NULL DEFAULT '',
  error TEXT
);
CREATE TABLE IF NOT EXISTS fragments (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  source_id TEXT NOT NULL REFERENCES documents(source_id) ON DELETE CASCADE,
  locator TEXT NOT NULL,
  fragment_kind TEXT NOT NULL,
  text TEXT NOT NULL,
  text_norm TEXT NOT NULL,
  hidden INTEGER NOT NULL DEFAULT 0,
  formula TEXT,
  cached_value TEXT,
  number_format TEXT,
  unit TEXT,
  base_date TEXT,
  meta TEXT NOT NULL DEFAULT '{}'
);
CREATE INDEX IF NOT EXISTS ix_fragments_source ON fragments(source_id);
CREATE INDEX IF NOT EXISTS ix_fragments_hidden ON fragments(hidden);
"""

_SELECT_DOCS = "SELECT source_id, path, root, kind, status, sha256, size, mtime, revision, title, author, created, modified, approval_state, doc_kind, synthetic, flags, n_fragments, hidden_fragments, indexed_at, error FROM documents"
_INSERT_DOC = "INSERT INTO documents(source_id, path, root, kind, status, sha256, size, mtime, revision, title, author, created, modified, approval_state, doc_kind, synthetic, flags, n_fragments, hidden_fragments, indexed_at, error) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"


def normalize_text(text: str) -> str:
    """검색용 정규화: 소문자 + 숫자 사이 콤마 제거 변형을 함께 저장한다."""
    low = text.lower()
    no_comma = re.sub(r"(?<=\d),(?=\d{3})", "", low)
    return low if no_comma == low else f"{low}\n{no_comma}"


def _row_to_doc(row: sqlite3.Row) -> DocumentMeta:
    return DocumentMeta(
        source_id=row["source_id"],
        path=row["path"],
        root=row["root"],
        kind=row["kind"],
        status=row["status"],
        sha256=row["sha256"],
        size=int(row["size"]),
        mtime=float(row["mtime"]),
        revision=row["revision"],
        title=row["title"],
        author=row["author"],
        created=row["created"],
        modified=row["modified"],
        approval_state=row["approval_state"],
        doc_kind=row["doc_kind"],
        synthetic=bool(row["synthetic"]),
        flags=json.loads(row["flags"]),
        n_fragments=int(row["n_fragments"]),
        hidden_fragments=int(row["hidden_fragments"]),
        indexed_at=row["indexed_at"],
        error=row["error"],
    )


def _row_to_fragment(row: sqlite3.Row) -> Fragment:
    return Fragment(
        locator=row["locator"],
        fragment_kind=row["fragment_kind"],
        text=row["text"],
        hidden=bool(row["hidden"]),
        formula=row["formula"],
        cached_value=row["cached_value"],
        number_format=row["number_format"],
        unit=row["unit"],
        base_date=row["base_date"],
        meta=json.loads(row["meta"]),
    )


class DocumentIndex:
    """sqlite 기반 로컬 문서 색인. embedding/OCR 없음."""

    def __init__(self, db_path: str | os.PathLike[str]) -> None:
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._conn = sqlite3.connect(str(self.db_path), timeout=30)
        self._conn.row_factory = sqlite3.Row
        self._conn.execute("PRAGMA foreign_keys=ON")
        self._conn.execute("PRAGMA busy_timeout=30000")
        self._conn.executescript(_SCHEMA)
        self._conn.execute(
            "INSERT OR IGNORE INTO meta(key, value) VALUES ('schema_version', ?)",
            (str(INDEX_SCHEMA_VERSION),),
        )
        self._conn.commit()

    # -- lifecycle ---------------------------------------------------------
    def close(self) -> None:
        self._conn.close()

    def __enter__(self) -> DocumentIndex:
        return self

    def __exit__(self, *exc: object) -> None:
        self.close()

    # -- 조회 --------------------------------------------------------------
    def documents(self) -> list[DocumentMeta]:
        rows = self._conn.execute(_SELECT_DOCS + " ORDER BY path").fetchall()
        return [_row_to_doc(r) for r in rows]

    def get_document(self, source_id: str) -> DocumentMeta | None:
        row = self._conn.execute(_SELECT_DOCS + " WHERE source_id=?", (source_id,)).fetchone()
        return _row_to_doc(row) if row else None

    def find_by_path(self, path: str | os.PathLike[str]) -> DocumentMeta | None:
        row = self._conn.execute(_SELECT_DOCS + " WHERE path=?", (str(Path(path).resolve()),)).fetchone()
        return _row_to_doc(row) if row else None

    def fragments(self, source_id: str, *, include_hidden: bool = True) -> list[Fragment]:
        sql = "SELECT * FROM fragments WHERE source_id=?"
        if not include_hidden:
            sql += " AND hidden=0"
        rows = self._conn.execute(sql + " ORDER BY id", (source_id,)).fetchall()
        return [_row_to_fragment(r) for r in rows]

    def candidate_fragments(
        self, patterns: list[str], *, include_hidden: bool = False, limit: int = 5000
    ) -> list[tuple[DocumentMeta, Fragment]]:
        """text_norm 에 patterns 중 하나라도 포함된 fragment (검색 1차 후보)."""
        if not patterns:
            return []
        # 값은 모두 바인딩 파라미터이며 patterns 개수만큼 'LIKE ?' 절을 반복한다 (문자열 삽입 없음)
        where = " OR ".join(["f.text_norm LIKE ? ESCAPE '\\'"] * len(patterns))
        params: list[Any] = [f"%{_like_escape(p)}%" for p in patterns]
        sql = "SELECT f.* FROM fragments f WHERE (" + where + ")"  # noqa: S608 - 바인딩 파라미터만 사용
        if not include_hidden:
            sql += " AND f.hidden=0"
        sql += " ORDER BY f.id LIMIT ?"
        params.append(limit)
        rows = self._conn.execute(sql, params).fetchall()
        docs: dict[str, DocumentMeta] = {}
        out: list[tuple[DocumentMeta, Fragment]] = []
        for r in rows:
            sid = r["source_id"]
            if sid not in docs:
                d = self.get_document(sid)
                if d is None:
                    continue
                docs[sid] = d
            out.append((docs[sid], _row_to_fragment(r)))
        return out

    def stats(self) -> dict[str, Any]:
        n_docs = self._conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        n_frag = self._conn.execute("SELECT COUNT(*) FROM fragments").fetchone()[0]
        n_hidden = self._conn.execute("SELECT COUNT(*) FROM fragments WHERE hidden=1").fetchone()[0]
        by_status = {
            r[0]: r[1]
            for r in self._conn.execute("SELECT status, COUNT(*) FROM documents GROUP BY status").fetchall()
        }
        return {
            "db_path": str(self.db_path),
            "schema_version": INDEX_SCHEMA_VERSION,
            "n_documents": int(n_docs),
            "n_fragments": int(n_frag),
            "n_hidden_fragments": int(n_hidden),
            "by_status": by_status,
        }

    # -- 변경 --------------------------------------------------------------
    def remove_document(self, source_id: str) -> None:
        with self._conn:
            self._conn.execute("DELETE FROM fragments WHERE source_id=?", (source_id,))
            self._conn.execute("DELETE FROM documents WHERE source_id=?", (source_id,))

    def prune_missing(self) -> list[str]:
        """디스크에 없거나 읽을 수 없는(권한 철회) 문서를 색인에서 제거한다."""
        removed: list[str] = []
        for d in self.documents():
            p = Path(d.path)
            if not p.is_file() or not os.access(p, os.R_OK):
                self.remove_document(d.source_id)
                removed.append(d.path)
        return removed

    def index_roots(
        self,
        roots: Iterable[str | os.PathLike[str]],
        *,
        include_hidden_flag: bool = True,
        include_notes: bool = True,
        max_file_mb: int = 200,
        prune: bool = True,
        forbid_links: bool = True,
    ) -> IndexReport:
        """허용 root 아래의 문서만 색인한다. root 밖은 절대 읽지 않는다.

        include_hidden_flag=True 이면 숨김 콘텐츠를 hidden=1 로 저장(검색 기본 제외), False 면 저장하지 않는다.
        """
        started = now_iso()
        root_paths = [Path(r).expanduser().resolve() for r in roots]
        for r in root_paths:
            if not r.is_dir():
                raise AgentError(
                    "E_PATH_OUTSIDE_ROOT", f"색인 root 폴더가 없습니다: {r.name}", details={"root": str(r)}
                )
        report = IndexReport(roots=[str(r) for r in root_paths], started_at=started, finished_at=started)
        seen: set[str] = set()
        for root in root_paths:
            for file_path in _walk_files(root, forbid_links=forbid_links):
                key = str(file_path)
                seen.add(key)
                meta = self.index_file(
                    file_path,
                    root,
                    include_hidden_flag=include_hidden_flag,
                    include_notes=include_notes,
                    max_file_mb=max_file_mb,
                )
                if meta.status == "INDEXED":
                    report.indexed.append(key)
                elif meta.status == "UNCHANGED":
                    report.unchanged.append(key)
                elif meta.status == "NEEDS_APPROVED_CONVERSION":
                    report.needs_conversion.append(key)
                elif meta.status == "SKIPPED_SIZE":
                    report.skipped.append({"path": key, "reason": meta.error or "size"})
                else:
                    report.errors.append({"path": key, "reason": meta.error or "unknown"})
                if meta.flags:
                    report.flagged[key] = meta.flags
        if prune:
            for d in self.documents():
                under_root = any(_is_under(Path(d.path), r) for r in root_paths)
                if under_root and d.path not in seen:
                    self.remove_document(d.source_id)
                    report.pruned.append(d.path)
            report.pruned.extend(self.prune_missing())
        st = self.stats()
        report.n_documents = st["n_documents"]
        report.n_fragments = st["n_fragments"]
        report.n_hidden_fragments = st["n_hidden_fragments"]
        report.finished_at = now_iso()
        return report

    def index_file(
        self,
        path: Path,
        root: Path,
        *,
        include_hidden_flag: bool = True,
        include_notes: bool = True,
        max_file_mb: int = 200,
    ) -> DocumentMeta:
        path = Path(path).resolve()
        root = Path(root).resolve()
        if not _is_under(path, root):
            raise AgentError("E_PATH_OUTSIDE_ROOT", f"색인 대상이 root 밖입니다: {path.name}")
        ext = path.suffix.lower()
        st = path.stat()
        source_id = stable_id("doc", str(path))
        existing = self.get_document(source_id)
        base = DocumentMeta(
            source_id=source_id,
            path=str(path),
            root=str(root),
            kind=SUPPORTED_EXT.get(ext) or LEGACY_EXT.get(ext) or ext.lstrip("."),
            status="ERROR",
            sha256="",
            size=int(st.st_size),
            mtime=float(st.st_mtime),
            revision=revision_from_filename(path.name) or "",
            title="",
            author="",
            indexed_at=now_iso(),
        )
        if ext in LEGACY_EXT:
            base.sha256 = sha256_file(path)
            base.status = "NEEDS_APPROVED_CONVERSION"
            base.error = "구형 형식(.ppt/.xls)은 파싱하지 않습니다. 승인된 변환 후 .pptx/.xlsx 로 색인하세요."
            self._upsert(base, [])
            return base
        if ext not in SUPPORTED_EXT:
            raise AgentError("E_DOC_UNSUPPORTED", f"지원되지 않는 문서 확장자: {ext}")
        if st.st_size > max_file_mb * 1024 * 1024:
            base.sha256 = ""
            base.status = "SKIPPED_SIZE"
            base.error = f"파일 크기 {st.st_size / (1024 * 1024):.1f} MB 가 상한 {max_file_mb} MB 를 초과"
            self._upsert(base, [])
            return base
        base.sha256 = sha256_file(path)
        if (
            existing is not None
            and existing.status == "INDEXED"
            and existing.sha256 == base.sha256
            and abs(existing.mtime - base.mtime) < 1e-6
        ):
            existing.status = "UNCHANGED"
            return existing
        try:
            base.flags = scan_package_flags(path)
            if base.kind == "pptx":
                frags, props = extract_pptx(path, include_notes=include_notes)
            else:
                frags, props = extract_xlsx(path)
        except AgentError as exc:
            base.status = "ERROR"
            base.error = f"{exc.code}: {exc.message}"[:500]
            self._upsert(base, [])
            return base
        except Exception as exc:  # noqa: BLE001 - 손상 파일 등은 오류 상태로 기록하고 계속
            base.status = "ERROR"
            base.error = f"{type(exc).__name__}: {str(exc)[:300]}"
            self._upsert(base, [])
            return base
        if not include_hidden_flag:
            frags = [f for f in frags if not f.hidden]
        base.title = str(props.get("title") or "")
        base.author = str(props.get("author") or "")
        base.created = props.get("created")
        base.modified = props.get("modified")
        base.revision = base.revision or str(props.get("revision") or "")
        base.approval_state = normalize_approval(
            str(props.get("content_status") or "") or _approval_from_name(path.name)
        )
        base.synthetic = _is_synthetic_props(
            props.get("category"), props.get("keywords"), props.get("comments"), props.get("title")
        )
        base.doc_kind = classify_doc_kind(path, root, frags)
        base.n_fragments = len(frags)
        base.hidden_fragments = sum(1 for f in frags if f.hidden)
        base.status = "INDEXED"
        base.error = None
        self._upsert(base, frags)
        return base

    def _upsert(self, doc: DocumentMeta, frags: list[Fragment]) -> None:
        with self._conn:
            self._conn.execute("DELETE FROM fragments WHERE source_id=?", (doc.source_id,))
            self._conn.execute("DELETE FROM documents WHERE source_id=?", (doc.source_id,))
            self._conn.execute(
                _INSERT_DOC,
                (
                    doc.source_id,
                    doc.path,
                    doc.root,
                    doc.kind,
                    doc.status,
                    doc.sha256,
                    doc.size,
                    doc.mtime,
                    doc.revision,
                    doc.title,
                    doc.author,
                    doc.created,
                    doc.modified,
                    doc.approval_state,
                    doc.doc_kind,
                    1 if doc.synthetic else 0,
                    json.dumps(doc.flags, ensure_ascii=False),
                    doc.n_fragments,
                    doc.hidden_fragments,
                    doc.indexed_at,
                    doc.error,
                ),
            )
            self._conn.executemany(
                "INSERT INTO fragments(source_id, locator, fragment_kind, text, text_norm, hidden, formula, cached_value, number_format, unit, base_date, meta) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                [
                    (
                        doc.source_id,
                        f.locator,
                        f.fragment_kind,
                        f.text,
                        normalize_text(f.text),
                        1 if f.hidden else 0,
                        f.formula,
                        f.cached_value,
                        f.number_format,
                        f.unit,
                        f.base_date,
                        json.dumps(f.meta, ensure_ascii=False),
                    )
                    for f in frags
                ],
            )


def _approval_from_name(name: str) -> str:
    low = name.lower()
    if "approved" in low or "승인" in low:
        return "approved"
    if "draft" in low or "초안" in low:
        return "draft"
    return ""


def _like_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _is_under(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _walk_files(root: Path, *, forbid_links: bool) -> Iterator[Path]:
    """root 아래의 지원 문서 파일을 결정적 순서로 순회. symlink/junction 은 따라가지 않는다."""
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        dirnames[:] = sorted(d for d in dirnames if not (forbid_links and Path(dirpath, d).is_symlink()))
        for fn in sorted(filenames):
            if fn.startswith("~$") or fn.startswith("."):
                continue
            ext = Path(fn).suffix.lower()
            if ext not in SUPPORTED_EXT and ext not in LEGACY_EXT:
                continue
            p = Path(dirpath) / fn
            if forbid_links and has_link_component(p, stop_at=root):
                continue
            if not p.is_file():
                continue
            yield p
