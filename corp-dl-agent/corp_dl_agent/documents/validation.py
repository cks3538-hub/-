"""생성 문서 검증 → validation_report.json

검사 항목
- 수치 일치: 문서에서 다시 읽은 값(placeholder 위치의 텍스트 / named range 셀 값) == report_payload 값
- 구조 보존: slide 수·순서·layout·shape 이름, sheet 이름·순서·상태, named range, 표, 수식 셀(수식 원문) 이 템플릿과 동일
- 오래된 값 잔존: 과거 차종/날짜/원가 등 stale_values 가 산출물 어디에도(숨김 slide/sheet 포함) 남지 않음
- placeholder 잔존: '{{' 가 남아 있지 않음
- 원본 template hash 불변, 산출물 hash 가 manifest 와 일치
- 상태 명시: RECALC_NOT_RUN(수식 재계산 안 함) / RENDER_NOT_RUN(렌더링 검사 안 함)

passed 는 FAIL 이 하나도 없을 때 True (NOT_RUN/PARTIAL 은 허용하되 보고서에 남긴다).
"""

from __future__ import annotations

import math
import os
import re
from pathlib import Path
from typing import Any, Literal

from pydantic import Field

from corp_dl_agent.common import Quantity, Status, StrictModel, atomic_write_json, now_iso, sha256_file
from corp_dl_agent.documents.payload import (
    CONFLICT_TEXT,
    MISSING_TEXT,
    PayloadValidation,
    ReportPayload,
    parse_number,
    validate_payload,
)
from corp_dl_agent.documents.templates import (
    DocumentManifest,
    RecalcStatus,
    RenderStatus,
    collect_pptx_texts,
    collect_xlsx_texts,
    defined_name_destinations,
    read_pptx_structure,
    read_xlsx_structure,
)
from corp_dl_agent.errors import blocked_dependency
from corp_dl_agent.version import SCHEMA_VERSION

Category = Literal["numeric", "text", "structure", "stale", "placeholder", "hash", "status", "payload"]
_PLACEHOLDER_LEFT_RE = re.compile(r"\{\{[^}]*\}\}")
_NUMERIC_TOKEN_RE = re.compile(r"^-?\d[\d,]*(?:\.\d+)?$")


class ValidationCheck(StrictModel):
    check_id: str
    category: Category
    status: Status
    locator: str | None = None
    expected: str | None = None
    actual: str | None = None
    message_ko: str = ""


class DocumentValidation(StrictModel):
    document_type: Literal["pptx", "xlsx"]
    path: str
    sha256: str
    template_path: str | None = None
    checks: list[ValidationCheck] = Field(default_factory=list)
    n_pass: int = 0
    n_fail: int = 0
    n_not_run: int = 0
    n_partial: int = 0
    passed: bool = True


class ValidationReport(StrictModel):
    schema_version: str = SCHEMA_VERSION
    created_at: str
    payload_report_id: str
    payload_passed: bool
    payload_checks: list[dict[str, Any]] = Field(default_factory=list)
    payload_provenance_issues: list[str] = Field(default_factory=list)
    n_missing: int = 0
    n_conflict: int = 0
    documents: list[DocumentValidation] = Field(default_factory=list)
    stale_values: list[str] = Field(default_factory=list)
    recalc_status: RecalcStatus = "RECALC_NOT_RUN"
    render_status: RenderStatus = "RENDER_NOT_RUN"
    renderer: str = "none"
    recalc_engine: str = "none"
    passed: bool
    synthetic: bool
    notes: list[str] = Field(default_factory=list)


# ---------------------------------------------------------------------------
# 보조
# ---------------------------------------------------------------------------


def stale_pattern(token: str) -> re.Pattern[str]:
    """숫자 토큰은 자릿수 경계·콤마 유무를 함께 검사, 문자열 토큰은 대소문자 무시 부분 일치."""
    t = token.strip()
    if _NUMERIC_TOKEN_RE.match(t):
        digits = t.replace(",", "")
        int_part, _, frac = digits.partition(".")
        grouped = f"{int(int_part):,}" if int_part.lstrip("-").isdigit() else int_part
        alts = {re.escape(digits), re.escape(grouped + (f".{frac}" if frac else ""))}
        return re.compile(r"(?<![\d.,])(?:" + "|".join(sorted(alts)) + r")(?![\d.,])")
    return re.compile(re.escape(t), re.IGNORECASE)


def find_stale(texts: dict[str, str], stale_values: list[str]) -> list[tuple[str, str]]:
    found: list[tuple[str, str]] = []
    pats = [(tok, stale_pattern(tok)) for tok in stale_values if tok.strip()]
    for loc, text in texts.items():
        for tok, pat in pats:
            if pat.search(text):
                found.append((loc, tok))
    return found


def _numbers_equal(a: float, b: float, *, rel: float = 1e-9, abs_tol: float = 1e-9) -> bool:
    return math.isclose(a, b, rel_tol=rel, abs_tol=abs_tol)


def _summarize(doc: DocumentValidation) -> DocumentValidation:
    doc.n_pass = sum(1 for c in doc.checks if c.status is Status.PASS)
    doc.n_fail = sum(1 for c in doc.checks if c.status is Status.FAIL)
    doc.n_not_run = sum(1 for c in doc.checks if c.status is Status.NOT_RUN)
    doc.n_partial = sum(1 for c in doc.checks if c.status is Status.PARTIAL)
    doc.passed = doc.n_fail == 0
    return doc


def _hash_checks(path: Path, manifest: DocumentManifest | None, prefix: str) -> list[ValidationCheck]:
    out: list[ValidationCheck] = []
    if manifest is None:
        out.append(
            ValidationCheck(
                check_id=f"{prefix}.template_hash",
                category="hash",
                status=Status.NOT_RUN,
                message_ko="manifest 가 없어 템플릿 hash 를 비교하지 못했습니다",
            )
        )
        return out
    tpl = Path(manifest.template_path)
    if tpl.is_file():
        actual = sha256_file(tpl)
        ok = actual == manifest.input_hash
        out.append(
            ValidationCheck(
                check_id=f"{prefix}.template_hash",
                category="hash",
                status=Status.PASS if ok else Status.FAIL,
                locator=str(tpl),
                expected=manifest.input_hash,
                actual=actual,
                message_ko="원본 템플릿 hash 불변" if ok else "원본 템플릿이 생성 이후 변경되었습니다",
            )
        )
    else:
        out.append(
            ValidationCheck(
                check_id=f"{prefix}.template_hash",
                category="hash",
                status=Status.NOT_RUN,
                locator=str(tpl),
                message_ko="템플릿 파일을 찾을 수 없습니다",
            )
        )
    actual_out = sha256_file(path)
    ok = actual_out == manifest.output_hash
    out.append(
        ValidationCheck(
            check_id=f"{prefix}.output_hash",
            category="hash",
            status=Status.PASS if ok else Status.FAIL,
            locator=str(path),
            expected=manifest.output_hash,
            actual=actual_out,
            message_ko="산출물 hash 가 manifest 와 일치"
            if ok
            else "산출물이 manifest 기록 이후 변경되었습니다",
        )
    )
    return out


def _stale_and_placeholder_checks(
    texts: dict[str, str], stale_values: list[str], prefix: str
) -> list[ValidationCheck]:
    out: list[ValidationCheck] = []
    leftovers = [loc for loc, t in texts.items() if _PLACEHOLDER_LEFT_RE.search(t)]
    out.append(
        ValidationCheck(
            check_id=f"{prefix}.placeholders",
            category="placeholder",
            status=Status.FAIL if leftovers else Status.PASS,
            locator=", ".join(leftovers[:10]) or None,
            message_ko=f"치환되지 않은 placeholder {len(leftovers)}곳"
            if leftovers
            else "남은 placeholder 없음",
        )
    )
    if stale_values:
        found = find_stale(texts, stale_values)
        out.append(
            ValidationCheck(
                check_id=f"{prefix}.stale_values",
                category="stale",
                status=Status.FAIL if found else Status.PASS,
                locator=", ".join(f"{loc}[{tok}]" for loc, tok in found[:10]) or None,
                expected="none",
                actual=", ".join(sorted({tok for _, tok in found})) or "none",
                message_ko=f"오래된 값 잔존 {len(found)}곳 (과거 차종/날짜/원가)"
                if found
                else "오래된 값 잔존 없음",
            )
        )
    else:
        out.append(
            ValidationCheck(
                check_id=f"{prefix}.stale_values",
                category="stale",
                status=Status.NOT_RUN,
                message_ko="stale_values 가 지정되지 않아 잔존 검사를 건너뜀",
            )
        )
    return out


def _conflict_after_revalidation(check_id: str, locator: str, rendered: str, notes: str | None) -> ValidationCheck:
    """문서에는 값이 기록되었으나 payload 재검증(합계/차이) 에서 해당 항목이 CONFLICT 로 바뀐 경우."""
    return ValidationCheck(
        check_id=check_id,
        category="numeric",
        status=Status.FAIL,
        locator=locator,
        expected=CONFLICT_TEXT,
        actual=rendered,
        message_ko=f"문서에 기록된 값이 payload 재검증에서 CONFLICT 로 판정됨: {notes or ''}".rstrip(": "),
    )


def _value_without_payload(check_id: str, locator: str, rendered: str) -> ValidationCheck:
    """문서에는 값이 있으나 payload 에는 값이 없는(MISSING) 경우 — 근거 없는 숫자."""
    return ValidationCheck(
        check_id=check_id,
        category="numeric",
        status=Status.FAIL,
        locator=locator,
        expected=MISSING_TEXT,
        actual=rendered,
        message_ko="payload 에 값이 없는데(MISSING) 문서에는 값이 기록되어 있습니다",
    )


def _engine_status_checks(prefix: str, *, renderer: str, recalc_engine: str | None) -> list[ValidationCheck]:
    """RECALC/RENDER 상태 항목. 엔진이 설정되어 있어도 이 버전은 실행하지 않는다 (E_NOT_SUPPORTED, 상태는 NOT_RUN)."""
    out: list[ValidationCheck] = []
    if recalc_engine is not None:
        msg = "RECALC_NOT_RUN: openpyxl 은 수식을 계산하지 않으며 저장된 수식 셀에는 cached value 가 없습니다. 수식 결과값은 Excel/승인 엔진 재계산 후 확인하세요"
        if recalc_engine != "none":
            msg += f" (documents.recalc_engine='{recalc_engine}' 설정됨 — 이 버전은 재계산 엔진 실행을 지원하지 않습니다: E_NOT_SUPPORTED)"
        out.append(
            ValidationCheck(
                check_id=f"{prefix}.recalc",
                category="status",
                status=Status.NOT_RUN,
                actual=f"recalc_engine={recalc_engine}",
                message_ko=msg,
            )
        )
    what = "slide 렌더링(잘림/겹침/폰트)" if prefix == "pptx" else "인쇄영역/렌더링"
    msg = f"RENDER_NOT_RUN: {what} 검사는 승인 렌더러가 없어 수행하지 않음"
    if renderer != "none":
        msg = f"RENDER_NOT_RUN: documents.renderer='{renderer}' 설정됨 — 이 버전은 렌더러 실행을 지원하지 않습니다 (E_NOT_SUPPORTED). {what} 검사는 수행하지 않음"
    out.append(
        ValidationCheck(
            check_id=f"{prefix}.render",
            category="status",
            status=Status.NOT_RUN,
            actual=f"renderer={renderer}",
            message_ko=msg,
        )
    )
    return out


def _structure_diff(expected: dict[str, Any], actual: dict[str, Any], keys: list[str]) -> list[str]:
    diffs: list[str] = []
    for k in keys:
        if expected.get(k) != actual.get(k):
            diffs.append(k)
    return diffs


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def validate_pptx(
    path: str | os.PathLike[str],
    payload: ReportPayload,
    *,
    manifest: DocumentManifest | None = None,
    stale_values: list[str] | None = None,
    renderer: str = "none",
) -> DocumentValidation:
    p = Path(path)
    doc = DocumentValidation(
        document_type="pptx",
        path=str(p),
        sha256=sha256_file(p),
        template_path=manifest.template_path if manifest else None,
    )
    texts = collect_pptx_texts(p, include_hidden=True)
    checks = doc.checks
    if manifest is None:
        checks.append(
            ValidationCheck(
                check_id="pptx.numeric",
                category="numeric",
                status=Status.NOT_RUN,
                message_ko="manifest 가 없어 placeholder 위치별 수치 대조를 건너뜀",
            )
        )
    else:
        n_ok = 0
        for r in manifest.replacements:
            text = texts.get(r.locator)
            if r.status == "unapproved":
                continue
            if text is None:
                checks.append(
                    ValidationCheck(
                        check_id=f"pptx.numeric[{r.locator}:{r.key}]",
                        category="numeric",
                        status=Status.FAIL,
                        locator=r.locator,
                        message_ko="locator 를 산출물에서 찾을 수 없습니다",
                    )
                )
                continue
            if r.status in ("missing", "unknown_key"):
                checks.append(
                    ValidationCheck(
                        check_id=f"pptx.missing[{r.locator}:{r.key}]",
                        category="numeric",
                        status=Status.PARTIAL if r.rendered in text else Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        message_ko="payload 에 값이 없어 MISSING 으로 표시됨"
                        if r.status == "missing"
                        else f"payload 에 없는 key '{r.key}' → MISSING 표시",
                    )
                )
                continue
            if r.status == "conflict":
                checks.append(
                    ValidationCheck(
                        check_id=f"pptx.conflict[{r.locator}:{r.key}]",
                        category="numeric",
                        status=Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        message_ko="출처 상충(CONFLICT) 값이 문서에 표시됨",
                    )
                )
                continue
            if r.rendered not in text:
                checks.append(
                    ValidationCheck(
                        check_id=f"pptx.numeric[{r.locator}:{r.key}]",
                        category="numeric" if r.numeric else "text",
                        status=Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        actual=text[:120],
                        message_ko="문서에서 렌더링된 값을 찾지 못했습니다",
                    )
                )
                continue
            if r.numeric:
                item = payload.items.get(r.key)
                pq = item.quantity if item else None
                if pq is not None and pq.value_type == "conflict":
                    checks.append(
                        _conflict_after_revalidation(
                            f"pptx.conflict[{r.locator}:{r.key}]", r.locator, r.rendered, pq.notes
                        )
                    )
                    continue
                pv = pq.value if pq else None
                if pv is None:
                    checks.append(
                        _value_without_payload(f"pptx.numeric[{r.locator}:{r.key}]", r.locator, r.rendered)
                    )
                    continue
                start = text.index(r.rendered)
                read_back = parse_number(text[start : start + len(r.rendered)])
                if read_back is None or not _numbers_equal(read_back, pv):
                    checks.append(
                        ValidationCheck(
                            check_id=f"pptx.numeric[{r.locator}:{r.key}]",
                            category="numeric",
                            status=Status.FAIL,
                            locator=r.locator,
                            expected=str(pv),
                            actual=str(read_back),
                            message_ko="문서에서 읽은 수치가 payload 와 다릅니다",
                        )
                    )
                    continue
            n_ok += 1
        checks.append(
            ValidationCheck(
                check_id="pptx.numeric",
                category="numeric",
                status=Status.PASS,
                message_ko=f"placeholder {n_ok}곳의 값이 payload 와 일치",
            )
            if not any(c.status is Status.FAIL and c.category in ("numeric", "text") for c in checks)
            else ValidationCheck(
                check_id="pptx.numeric",
                category="numeric",
                status=Status.FAIL,
                message_ko="일부 placeholder 값이 payload 와 다릅니다",
            )
        )
    checks.extend(_stale_and_placeholder_checks(texts, list(stale_values or payload.stale_values), "pptx"))
    # 구조
    if manifest is not None and Path(manifest.template_path).is_file():
        tpl = read_pptx_structure(manifest.template_path)
        cur = read_pptx_structure(p)
        keys = ["n_slides", "slide_ids", "n_masters", "n_layouts", "n_charts", "n_tables"]
        diffs = _structure_diff(tpl, cur, keys)
        tpl_slides = [(s["slide_id"], s["layout"], s["hidden"], s["shapes"]) for s in tpl["slides"]]
        cur_slides = [(s["slide_id"], s["layout"], s["hidden"], s["shapes"]) for s in cur["slides"]]
        if tpl_slides != cur_slides:
            diffs.append("slides(layout/hidden/shapes)")
        checks.append(
            ValidationCheck(
                check_id="pptx.structure",
                category="structure",
                status=Status.FAIL if diffs else Status.PASS,
                expected=f"slides={tpl['n_slides']}, charts={tpl['n_charts']}, tables={tpl['n_tables']}",
                actual=f"slides={cur['n_slides']}, charts={cur['n_charts']}, tables={cur['n_tables']}",
                message_ko=("구조 차이: " + ", ".join(diffs))
                if diffs
                else "slide 수·순서·layout·shape·차트·표 보존",
            )
        )
    else:
        checks.append(
            ValidationCheck(
                check_id="pptx.structure",
                category="structure",
                status=Status.NOT_RUN,
                message_ko="템플릿을 찾을 수 없어 구조 비교를 건너뜀",
            )
        )
    checks.extend(_hash_checks(p, manifest, "pptx"))
    checks.extend(_engine_status_checks("pptx", renderer=renderer, recalc_engine=None))
    return _summarize(doc)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def read_xlsx_values(path: str | os.PathLike[str]) -> tuple[dict[str, Any], dict[str, str]]:
    """(locator -> 원시 값, locator -> 수식) . range:<name> 단일 셀과 range:tbl_x/r{i}c{j} 도 포함."""
    try:
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("openpyxl", "documents.validation") from exc
    wb = openpyxl.load_workbook(str(path), data_only=False)
    values: dict[str, Any] = {}
    formulas: dict[str, str] = {}
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    loc = f"sheet:{ws.title}!{cell.coordinate}"
                    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                        formulas[loc] = str(cell.value)
                    else:
                        values[loc] = cell.value
        names: list[tuple[str, Any]] = [(str(k), v) for k, v in wb.defined_names.items()]
        for ws in wb.worksheets:
            names.extend((str(k), v) for k, v in ws.defined_names.items())
        for name, dn in names:
            for sheet, coord in defined_name_destinations(dn):
                if sheet not in wb.sheetnames:
                    continue
                c = coord.replace("$", "")
                ws = wb[sheet]
                if ":" not in c:
                    cell = ws[c]
                    if cell.data_type == "f" or (
                        isinstance(cell.value, str) and str(cell.value).startswith("=")
                    ):
                        formulas[f"range:{name}"] = str(cell.value)
                    else:
                        values[f"range:{name}"] = cell.value
                    continue
                min_col, min_row, max_col, max_row = range_boundaries(c)
                c0, r0, c1, r1 = int(min_col or 1), int(min_row or 1), int(max_col or 1), int(max_row or 1)
                for r in range(r0, r1 + 1):
                    for cc in range(c0, c1 + 1):
                        cell = ws.cell(row=r, column=cc)
                        loc = f"range:{name}/r{r - r0 + 1}c{cc - c0 + 1}"
                        if cell.data_type == "f":
                            formulas[loc] = str(cell.value)
                        else:
                            values[loc] = cell.value
    finally:
        wb.close()
    return values, formulas


def validate_xlsx(
    path: str | os.PathLike[str],
    payload: ReportPayload,
    *,
    manifest: DocumentManifest | None = None,
    stale_values: list[str] | None = None,
    renderer: str = "none",
    recalc_engine: str = "none",
) -> DocumentValidation:
    p = Path(path)
    doc = DocumentValidation(
        document_type="xlsx",
        path=str(p),
        sha256=sha256_file(p),
        template_path=manifest.template_path if manifest else None,
    )
    checks = doc.checks
    values, formulas = read_xlsx_values(p)
    if manifest is None:
        checks.append(
            ValidationCheck(
                check_id="xlsx.numeric",
                category="numeric",
                status=Status.NOT_RUN,
                message_ko="manifest 가 없어 named range 별 수치 대조를 건너뜀",
            )
        )
    else:
        n_ok = 0
        for r in manifest.replacements:
            if r.status == "unapproved":
                continue
            if r.locator in formulas:
                checks.append(
                    ValidationCheck(
                        check_id=f"xlsx.numeric[{r.locator}]",
                        category="numeric",
                        status=Status.FAIL,
                        locator=r.locator,
                        message_ko="값을 기록한 셀이 수식 셀로 바뀌었습니다",
                    )
                )
                continue
            if r.locator not in values:
                checks.append(
                    ValidationCheck(
                        check_id=f"xlsx.numeric[{r.locator}]",
                        category="numeric",
                        status=Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        message_ko="산출물에서 named range/셀을 찾지 못했습니다",
                    )
                )
                continue
            actual = values[r.locator]
            if r.status in ("missing", "unknown_key"):
                checks.append(
                    ValidationCheck(
                        check_id=f"xlsx.missing[{r.locator}]",
                        category="numeric",
                        status=Status.PARTIAL if str(actual) == r.rendered else Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        actual=str(actual),
                        message_ko="payload 에 값이 없어 MISSING 으로 표시됨",
                    )
                )
                continue
            if r.status == "conflict":
                checks.append(
                    ValidationCheck(
                        check_id=f"xlsx.conflict[{r.locator}]",
                        category="numeric",
                        status=Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        actual=str(actual),
                        message_ko="출처 상충(CONFLICT) 값이 문서에 표시됨",
                    )
                )
                continue
            if r.numeric and r.value is not None:
                # payload 재검증 이후의 값과 대조한다 (항목 또는 표 셀)
                pq: Quantity | None = None
                if r.key in payload.items and r.attr in (None, "value"):
                    pq = payload.items[r.key].quantity
                elif r.key in payload.tables and r.attr is not None:
                    pq = _table_cell_quantity(payload, r.key, r.attr, r.locator)
                if pq is not None and pq.value_type == "conflict":
                    checks.append(
                        _conflict_after_revalidation(f"xlsx.conflict[{r.locator}]", r.locator, r.rendered, pq.notes)
                    )
                    continue
                if pq is None or pq.value is None:
                    checks.append(_value_without_payload(f"xlsx.numeric[{r.locator}]", r.locator, r.rendered))
                    continue
                ok = (
                    isinstance(actual, (int, float))
                    and not isinstance(actual, bool)
                    and _numbers_equal(float(actual), float(r.value))
                    and _numbers_equal(float(actual), float(pq.value))
                )
                if not ok:
                    checks.append(
                        ValidationCheck(
                            check_id=f"xlsx.numeric[{r.locator}]",
                            category="numeric",
                            status=Status.FAIL,
                            locator=r.locator,
                            expected=str(r.value),
                            actual=str(actual),
                            message_ko="셀에서 읽은 수치가 payload 와 다릅니다",
                        )
                    )
                    continue
            elif str(actual) != r.rendered:
                checks.append(
                    ValidationCheck(
                        check_id=f"xlsx.text[{r.locator}]",
                        category="text",
                        status=Status.FAIL,
                        locator=r.locator,
                        expected=r.rendered,
                        actual=str(actual),
                        message_ko="셀 텍스트가 payload 와 다릅니다",
                    )
                )
                continue
            n_ok += 1
        failed = any(c.status is Status.FAIL and c.category in ("numeric", "text") for c in checks)
        checks.append(
            ValidationCheck(
                check_id="xlsx.numeric",
                category="numeric",
                status=Status.FAIL if failed else Status.PASS,
                message_ko="일부 named range 값이 payload 와 다릅니다"
                if failed
                else f"named range {n_ok}곳의 값이 payload 와 일치",
            )
        )
    texts = collect_xlsx_texts(p, include_hidden=True)
    checks.extend(_stale_and_placeholder_checks(texts, list(stale_values or payload.stale_values), "xlsx"))
    if manifest is not None and Path(manifest.template_path).is_file():
        tpl = read_xlsx_structure(manifest.template_path)
        cur = read_xlsx_structure(p)
        diffs = _structure_diff(tpl, cur, ["sheets", "sheet_states", "defined_names", "tables", "merged"])
        changed_formulas = sorted(
            k
            for k in set(tpl["formulas"]) | set(cur["formulas"])
            if tpl["formulas"].get(k) != cur["formulas"].get(k)
        )
        if changed_formulas:
            diffs.append(f"formulas({', '.join(changed_formulas[:10])})")
        checks.append(
            ValidationCheck(
                check_id="xlsx.structure",
                category="structure",
                status=Status.FAIL if diffs else Status.PASS,
                expected=f"sheets={tpl['sheets']}, names={len(tpl['defined_names'])}, formulas={len(tpl['formulas'])}",
                actual=f"sheets={cur['sheets']}, names={len(cur['defined_names'])}, formulas={len(cur['formulas'])}",
                message_ko=("구조 차이: " + ", ".join(diffs)) if diffs else "sheet·named range·표·수식 보존",
            )
        )
    else:
        checks.append(
            ValidationCheck(
                check_id="xlsx.structure",
                category="structure",
                status=Status.NOT_RUN,
                message_ko="템플릿을 찾을 수 없어 구조 비교를 건너뜀",
            )
        )
    checks.extend(_hash_checks(p, manifest, "xlsx"))
    engine_checks = _engine_status_checks("xlsx", renderer=renderer, recalc_engine=recalc_engine)
    n_formula_cells = sum(1 for k in formulas if k.startswith("sheet:"))
    engine_checks[0].actual = f"formula_cells={n_formula_cells}, recalc_engine={recalc_engine}"
    checks.extend(engine_checks)
    return _summarize(doc)


def _table_cell_quantity(payload: ReportPayload, table: str, column: str, locator: str) -> Quantity | None:
    """locator 'range:tbl_x/r{i}c{j}' 가 가리키는 payload 표 셀의 Quantity (없으면 None)."""
    m = re.search(r"/r(\d+)c(\d+)$", locator)
    if not m:
        return None
    row_idx = int(m.group(1)) - 1
    tb = payload.tables.get(table)
    if tb is None or row_idx >= len(tb.rows):
        return None
    return tb.rows[row_idx].cells.get(column)


# ---------------------------------------------------------------------------
# 통합
# ---------------------------------------------------------------------------


def validate_documents(
    pptx: str | os.PathLike[str] | None,
    xlsx: str | os.PathLike[str] | None,
    payload: ReportPayload,
    *,
    manifests: list[DocumentManifest] | None = None,
    stale_values: list[str] | None = None,
    payload_validation: PayloadValidation | None = None,
    renderer: str = "none",
    recalc_engine: str = "none",
) -> ValidationReport:
    pv = payload_validation or validate_payload(payload)
    stale = list(stale_values if stale_values is not None else payload.stale_values)
    by_type: dict[str, DocumentManifest] = {m.document_type: m for m in (manifests or [])}
    docs: list[DocumentValidation] = []
    if pptx is not None:
        docs.append(
            validate_pptx(pptx, pv.payload, manifest=by_type.get("pptx"), stale_values=stale, renderer=renderer)
        )
    if xlsx is not None:
        docs.append(
            validate_xlsx(
                xlsx,
                pv.payload,
                manifest=by_type.get("xlsx"),
                stale_values=stale,
                renderer=renderer,
                recalc_engine=recalc_engine,
            )
        )
    notes = [
        "수식 재계산(RECALC) 과 렌더링(RENDER) 은 수행하지 않았습니다. 상태는 각 문서의 status 항목에 있습니다.",
    ]
    if renderer != "none":
        notes.append(
            f"documents.renderer='{renderer}' 가 설정되어 있으나 이 버전은 렌더러 실행을 지원하지 않습니다 (E_NOT_SUPPORTED): RENDER_NOT_RUN."
        )
    if recalc_engine != "none":
        notes.append(
            f"documents.recalc_engine='{recalc_engine}' 가 설정되어 있으나 이 버전은 재계산 엔진 실행을 지원하지 않습니다 (E_NOT_SUPPORTED): RECALC_NOT_RUN."
        )
    if not pv.passed:
        notes.append("report_payload 재검증 실패: 합계/차이 불일치 또는 출처 누락 항목이 있습니다.")
    passed = pv.passed and all(d.passed for d in docs)
    return ValidationReport(
        created_at=now_iso(),
        payload_report_id=payload.report_id,
        payload_passed=pv.passed,
        payload_checks=[c.model_dump(mode="json") for c in pv.checks],
        payload_provenance_issues=pv.provenance_issues,
        n_missing=pv.n_missing,
        n_conflict=pv.n_conflict,
        documents=docs,
        stale_values=stale,
        renderer=renderer,
        recalc_engine=recalc_engine,
        passed=passed,
        synthetic=payload.synthetic,
        notes=notes,
    )


def write_validation_report(report: ValidationReport, path: str | os.PathLike[str]) -> Path:
    p = Path(path)
    atomic_write_json(p, report.model_dump(mode="json"))
    return p
