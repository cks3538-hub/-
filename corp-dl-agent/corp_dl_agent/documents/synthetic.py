"""합성(synthetic) 문서 fixture 생성기. 실제 회사 문서·차종·원가와 무관한 가상 자료다.

생성물 (make_synthetic_documents(dir)):
- past_review_2023.pptx      가상 과거 검토서: 구 차종 "X-OLD", 2023-05-01, 원가 1,234 원, 결론 문구, 표, 발표자 노트,
                              차트(계열 원데이터), 숨김 slide 1개
- past_cost_table.xlsx       가상 원가표: 수식(SUM) + 별도 cached value(일부는 의도적으로 오래된 캐시), named range,
                              ListObject 표, 숨김 시트 1개, 숨김 행 1개
- template_review.pptx       {{key}} placeholder 템플릿 (제목/본문/표 셀, run 이 나뉜 placeholder 포함)
- template_comparison.xlsx   named range 입력 셀 + 합계/차이 수식 유지 템플릿, tbl_cost_lines 표 범위
- report_payload.example.json  위 템플릿과 짝이 되는 예시 payload (새 값: X-NEW, 2026-09-01, 1,980 원)
- legacy_note.ppt            구형 확장자 자리표시 파일 (실제 바이너리 PPT 가 아님; NEEDS_APPROVED_CONVERSION 상태 시험용)
- README.md                  synthetic 표시
"""

from __future__ import annotations

import datetime as _dt
import os
import re
import zipfile
from pathlib import Path

from corp_dl_agent.common import Quantity, atomic_write_bytes, atomic_write_text
from corp_dl_agent.documents.payload import (
    PayloadCheck,
    PayloadColumn,
    PayloadItem,
    PayloadRow,
    PayloadTable,
    PayloadText,
    ReportPayload,
    write_payload,
)
from corp_dl_agent.errors import blocked_dependency

SYNTHETIC_TAG = "synthetic"
OLD_VEHICLE = "X-OLD"
OLD_DATE = "2023-05-01"
OLD_COST_TEXT = "1,234"
OLD_COST = 1234
NEW_VEHICLE = "X-NEW"
NEW_DATE = "2026-09-01"

_FIXED_DT = _dt.datetime(2023, 5, 1, 9, 0, 0)

README_TEXT = """# fixtures/documents (synthetic: true)

이 폴더의 모든 문서는 `corp_dl_agent.documents.synthetic.make_synthetic_documents` 가 생성한 **합성(synthetic)** 자료입니다.
실제 회사 검토서·원가표·양식을 이름만 바꾼 것이 아니며, 차종명(X-OLD/X-NEW)·날짜·원가·중량은 모두 가상 값입니다.

| 파일 | 내용 | 용도 |
|---|---|---|
| past_review_2023.pptx | 구 차종 X-OLD, 2023-05-01, 원가 1,234 원, 결론 문구, 표, 발표자 노트, 차트, 숨김 slide 1개 | 색인·근거 검색·오래된 값 회귀 시험 |
| past_cost_table.xlsx | 수식(SUM) + cached value(B9 는 의도적으로 오래된 캐시 120, 실제 123.4), named range(UnitCost/BaseDate/Vehicle), 표 T1, 숨김 시트·숨김 행 | 수식/cached 구분, hidden 제외 시험 |
| template_review.pptx | {{key}} placeholder (제목/본문/표 셀, run 이 나뉜 placeholder) | fill_pptx |
| template_comparison.xlsx | named range 입력 셀, 차이/합계 수식 유지, tbl_cost_lines 표 범위 | fill_xlsx |
| report_payload.example.json | 템플릿과 짝이 되는 예시 payload (X-NEW, 2026-09-01, 원가 1,980 원) | docs generate |
| legacy_note.ppt | 구형 확장자 자리표시(실제 PPT 바이너리 아님) | NEEDS_APPROVED_CONVERSION 상태 시험 |

core properties 의 category/keywords 에 `synthetic` 을 표시했습니다. 이 자료로 만든 산출물에는 `synthetic: true` 가 붙습니다.
"""


def _require_libs() -> None:
    try:
        import openpyxl  # noqa: F401
        import pptx  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise blocked_dependency("python-pptx/openpyxl", "documents.synthetic") from exc


def _set_pptx_props(prs: object, title: str, revision: int, status: str) -> None:
    cp = prs.core_properties  # type: ignore[attr-defined]
    cp.title = title
    cp.author = "corp-dl-agent synthetic"
    cp.last_modified_by = "corp-dl-agent synthetic"
    cp.category = SYNTHETIC_TAG
    cp.keywords = f"{SYNTHETIC_TAG}; corp-dl-agent; fixture"
    cp.comments = "합성 자료(synthetic). 실제 회사 문서가 아닙니다."
    cp.subject = "synthetic fixture"
    cp.content_status = status
    cp.revision = revision
    cp.created = _FIXED_DT
    cp.modified = _FIXED_DT
    cp.last_printed = _FIXED_DT


def _add_textbox(
    slide: object, name: str, left: float, top: float, width: float, height: float, lines: list[str]
) -> object:
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))  # type: ignore[attr-defined]
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(16)
    return box


def make_past_review_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = f"{OLD_VEHICLE} 설계 검토서"
    s1.placeholders[1].text = f"검토일 {OLD_DATE} / 합성 자료(synthetic)"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "원가·중량 요약"
    _add_textbox(
        s2,
        "Body 1",
        0.6,
        1.4,
        8.5,
        1.6,
        [
            f"부품 단가: {OLD_COST_TEXT} 원 ({OLD_DATE} 견적 기준)",
            "중량: 12.5 kg",
            f"결론: {OLD_VEHICLE} 설계안을 채택한다.",
        ],
    )
    shape = s2.shapes.add_table(4, 3, Inches(0.6), Inches(3.2), Inches(8.5), Inches(1.6))
    shape.name = "Table 1"
    tbl = shape.table
    cells = [
        ("항목", "값", "단위"),
        ("단가", OLD_COST_TEXT, "원"),
        ("중량", "12.5", "kg"),
        ("재료비", "500", "원"),
    ]
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = text
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(14)
    s2.notes_slide.notes_text_frame.text = (
        f"발표자 노트: 단가 {OLD_COST_TEXT} 원은 {OLD_DATE} 견적 기준. 합성 자료."
    )

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = f"중량 비교 ({OLD_VEHICLE} vs {OLD_VEHICLE}-B)"
    cd = CategoryChartData()
    cd.categories = [OLD_VEHICLE, f"{OLD_VEHICLE}-B"]
    cd.add_series("중량(kg)", (12.5, 11.8))
    gf = s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(7), Inches(4), cd)
    gf.name = "Chart 1"
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "중량 비교"

    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "내부 검토 메모 (숨김)"
    _add_textbox(s4, "Memo 1", 0.6, 1.4, 8.5, 1.2, ["구 단가 협의 내역: 1,100 원 제안 (미확정, 숨김 slide)"])
    s4._element.set("show", "0")

    _set_pptx_props(prs, f"{OLD_VEHICLE} 설계 검토서 (synthetic)", 3, "Approved")
    prs.save(str(path))
    return path


def make_template_review_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "{{subject}}"
    s1.placeholders[1].text = "{{report_id}} / 작성일 {{created_at}} / 출처 {{data_origin}}"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "설계안 요약"
    box = _add_textbox(
        s2,
        "Body 1",
        0.6,
        1.4,
        8.5,
        3.0,
        [
            "차종: {{vehicle}}",
            "기준일: {{base_date}}",
            "중량 A: {{mass_a}}",
            "중량 B: {{mass_b}}",
            "중량 차이(B-A): {{mass_delta}}",
        ],
    )
    tf = box.text_frame  # type: ignore[attr-defined]
    # run 이 나뉜 placeholder (PowerPoint 편집 중 흔히 생김)
    para = tf.add_paragraph()
    r1 = para.add_run()
    r1.text = "원가 합계: {{cost_"
    r1.font.bold = True
    r1.font.size = Pt(16)
    r2 = para.add_run()
    r2.text = "total}} (출처 {{cost_total.source}})"
    r2.font.size = Pt(16)
    para = tf.add_paragraph()
    run = para.add_run()
    run.text = "결론: {{conclusion}}"
    run.font.size = Pt(16)

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "비교표"
    shape = s3.shapes.add_table(4, 4, Inches(0.6), Inches(1.6), Inches(8.5), Inches(2.4))
    shape.name = "Table 1"
    tbl = shape.table
    cells = [
        ("항목", "A", "B", "차이(B-A)"),
        (
            "{{mass_a.label}}",
            "{{mass_a.value}} {{mass_a.unit}}",
            "{{mass_b.value}} {{mass_b.unit}}",
            "{{mass_delta.value}}",
        ),
        ("원가", "{{cost_a}}", "{{cost_b}}", "{{cost_delta}}"),
        ("예측 삽입력", "{{insertion_force_a}}", "{{insertion_force_b}}", "{{insertion_force_delta}}"),
    ]
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                for rr in p.runs:
                    rr.font.size = Pt(14)
    s3.notes_slide.notes_text_frame.text = "근거 요약: {{evidence_note}} (synthetic 템플릿)"

    _set_pptx_props(prs, "설계 검토서 템플릿 (synthetic)", 1, "Draft")
    prs.save(str(path))
    return path


def _inject_cached_values(path: Path, sheet_xml: str, cached: dict[str, str]) -> None:
    """openpyxl 은 수식 셀에 cached value 를 쓰지 않으므로 sheet XML 의 <f> 뒤에 <v> 를 넣는다 (fixture 전용)."""
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        data = {n: zf.read(n) for n in names}
        infos = {n: zf.getinfo(n) for n in names}
    xml = data[sheet_xml].decode("utf-8")
    for coord, value in cached.items():
        pat = re.compile(rf'(<c r="{coord}"[^>]*>)(<f>[^<]*</f>)(?:<v></v>|<v>[^<]*</v>)?')
        xml, n = pat.subn(rf"\1\2<v>{value}</v>", xml, count=1)
        if n != 1:
            raise RuntimeError(f"cached value 주입 실패: {coord}")
    data[sheet_xml] = xml.encode("utf-8")
    tmp = path.with_name(f".{path.name}.tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for member in names:
            info = zipfile.ZipInfo(member, date_time=infos[member].date_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            zo.writestr(info, data[member])
    os.replace(tmp, path)


def _set_xlsx_props(wb: object, title: str, revision: str, status: str) -> None:
    p = wb.properties  # type: ignore[attr-defined]
    p.title = title
    p.creator = "corp-dl-agent synthetic"
    p.lastModifiedBy = "corp-dl-agent synthetic"
    p.category = SYNTHETIC_TAG
    p.keywords = f"{SYNTHETIC_TAG}; corp-dl-agent; fixture"
    p.description = "합성 자료(synthetic). 실제 회사 문서가 아닙니다."
    p.contentStatus = status
    p.revision = revision
    p.created = _FIXED_DT
    p.modified = _FIXED_DT


def make_past_cost_xlsx(path: Path) -> Path:
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Cost"
    ws["A1"], ws["B1"], ws["C1"] = "항목", "금액(원)", "비고"
    lines = [
        ("재료비", 500, "PP 수지"),
        ("가공비", 400, "사출"),
        ("금형비", 200, "상각"),
        ("후처리", 100, "도장"),
        ("조립", 34, "수작업"),
    ]
    for i, (name, amount, note) in enumerate(lines, start=2):
        ws.cell(i, 1, name)
        ws.cell(i, 2, amount)
        ws.cell(i, 3, note)
    ws["A7"], ws["B7"] = "합계", "=SUM(B2:B6)"
    ws["B7"].number_format = "#,##0"
    ws["A8"], ws["B8"] = "기준일", OLD_DATE
    ws["A9"], ws["B9"] = "관리비(10%)", "=B7*0.1"
    ws["A10"], ws["B10"] = "차종", OLD_VEHICLE
    ws["A11"], ws["B11"] = "결론", f"{OLD_VEHICLE} 단가 {OLD_COST_TEXT} 원 기준 승인"
    ws.row_dimensions[9].hidden = True
    ws.add_table(Table(displayName="T1", ref="A1:C6"))
    wb.defined_names["UnitCost"] = DefinedName("UnitCost", attr_text="Cost!$B$7")
    wb.defined_names["BaseDate"] = DefinedName("BaseDate", attr_text="Cost!$B$8")
    wb.defined_names["Vehicle"] = DefinedName("Vehicle", attr_text="Cost!$B$10")
    memo = wb.create_sheet("내부메모")
    memo["A1"] = "숨김 시트: 구 단가 협의 1,100 원 (미확정)"
    memo.sheet_state = "hidden"
    _set_xlsx_props(wb, f"{OLD_VEHICLE} 원가표 (synthetic)", "2", "Approved")
    wb.save(str(path))
    # B7 은 실제 합계(1234) 캐시, B9 는 의도적으로 오래된 캐시(120; 실제 123.4) → 수식/cached 구분 시험
    _inject_cached_values(path, "xl/worksheets/sheet1.xml", {"B7": str(OLD_COST), "B9": "120"})
    return path


def make_template_comparison_xlsx(path: Path) -> Path:
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Comparison"
    header = [("A1", "제목"), ("A2", "보고서 ID"), ("A3", "작성일"), ("A4", "차종"), ("A5", "기준일")]
    for coord, text in header:
        ws[coord] = text
    names: dict[str, str] = {
        "subject": "$B$1",
        "report_id": "$B$2",
        "created_at": "$B$3",
        "vehicle": "$B$4",
        "base_date": "$B$5",
        "mass_a": "$B$8",
        "mass_b": "$C$8",
        "mass_a__unit": "$E$8",
        "cost_a": "$B$9",
        "cost_b": "$C$9",
        "cost_a__unit": "$E$9",
        "insertion_force_a": "$B$10",
        "insertion_force_b": "$C$10",
        "insertion_force_a__unit": "$E$10",
        "cost_line_material": "$B$13",
        "cost_line_process": "$B$14",
        "cost_line_tooling": "$B$15",
        "cost_total_formula": "$B$16",
        "cost_total": "$B$17",
        "data_origin": "$B$20",
        "conclusion": "$B$21",
    }
    ws["A7"], ws["B7"], ws["C7"], ws["D7"], ws["E7"] = "항목", "A", "B", "차이(B-A)", "단위"
    ws["A8"], ws["D8"] = "중량", "=C8-B8"
    ws["A9"], ws["D9"] = "원가", "=C9-B9"
    ws["A10"], ws["D10"] = "예측 삽입력", "=C10-B10"
    ws["A12"], ws["B12"] = "원가 항목(B안)", "금액(원)"
    ws["A13"], ws["A14"], ws["A15"] = "재료비", "가공비", "금형비"
    ws["A16"], ws["B16"] = "합계(수식)", "=SUM(B13:B15)"
    ws["A17"] = "합계(payload)"
    ws["A18"], ws["B18"] = "검증 차이(payload-수식)", "=B17-B16"
    ws["A20"] = "데이터 출처"
    ws["A21"] = "결론"
    for c in ("B8", "C8", "B9", "C9", "B13", "B14", "B15", "B16", "B17", "D9", "B18"):
        ws[c].number_format = "#,##0.##"
    for name, ref in names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=f"Comparison!{ref}")
    lines = wb.create_sheet("Lines")
    lines["A1"], lines["B1"] = "원가 항목", "금액(원)"
    lines["A8"], lines["B8"] = "합계(수식)", "=SUM(B2:B6)"
    wb.defined_names["tbl_cost_lines"] = DefinedName("tbl_cost_lines", attr_text="Lines!$A$2:$B$6")
    _set_xlsx_props(wb, "설계안 비교 템플릿 (synthetic)", "1", "Draft")
    wb.save(str(path))
    return path


def example_payload(*, created_at: str = "2026-09-17T00:00:00+00:00") -> ReportPayload:
    """템플릿과 짝이 되는 예시 payload. 모든 수치는 가상이며 출처 표시를 갖는다."""

    def src(v: float, unit: str, locator: str) -> Quantity:
        return Quantity(value=v, unit=unit, value_type="source", source_locator=locator)

    def calc(v: float, unit: str, cid: str) -> Quantity:
        return Quantity(
            value=v,
            unit=unit,
            value_type="calculated",
            calculation_id=cid,
            calculation_version="mass-cost-4.0",
        )

    def pred(v: float, unit: str, run: str) -> Quantity:
        return Quantity(value=v, unit=unit, value_type="predicted", model_run_id=run)

    items = {
        "mass_a": PayloadItem(key="mass_a", label_ko="중량 A", quantity=calc(11.2, "kg", "mass:asm_a:total")),
        "mass_b": PayloadItem(key="mass_b", label_ko="중량 B", quantity=calc(10.4, "kg", "mass:asm_b:total")),
        "mass_delta": PayloadItem(
            key="mass_delta", label_ko="중량 차이(B-A)", quantity=calc(-0.8, "kg", "compare:mass_delta")
        ),
        "cost_a": PayloadItem(
            key="cost_a", label_ko="원가 A", quantity=calc(2050, "원", "cost:asm_a:unit_cost")
        ),
        "cost_b": PayloadItem(
            key="cost_b", label_ko="원가 B", quantity=calc(1980, "원", "cost:asm_b:unit_cost")
        ),
        "cost_delta": PayloadItem(
            key="cost_delta", label_ko="원가 차이(B-A)", quantity=calc(-70, "원", "compare:cost_delta")
        ),
        "cost_line_material": PayloadItem(
            key="cost_line_material",
            label_ko="재료비(B)",
            quantity=src(900, "원", "cost_recipes.json#asm_b.material"),
        ),
        "cost_line_process": PayloadItem(
            key="cost_line_process", label_ko="가공비(B)", quantity=calc(780, "원", "cost:asm_b:process")
        ),
        "cost_line_tooling": PayloadItem(
            key="cost_line_tooling", label_ko="금형비(B)", quantity=calc(300, "원", "cost:asm_b:tooling")
        ),
        "cost_total": PayloadItem(
            key="cost_total", label_ko="원가 합계(B)", quantity=calc(1980, "원", "cost:asm_b:total")
        ),
        "insertion_force_a": PayloadItem(
            key="insertion_force_a",
            label_ko="예측 삽입력 A",
            quantity=pred(42.5, "N", "run-synthetic-reg-0001"),
        ),
        "insertion_force_b": PayloadItem(
            key="insertion_force_b",
            label_ko="예측 삽입력 B",
            quantity=pred(39.8, "N", "run-synthetic-reg-0001"),
        ),
        "insertion_force_delta": PayloadItem(
            key="insertion_force_delta",
            label_ko="삽입력 차이(B-A)",
            quantity=calc(-2.7, "N", "compare:insertion_force_delta"),
        ),
        "retention_pass_rate": PayloadItem(
            key="retention_pass_rate",
            label_ko="유지력 합격률",
            quantity=Quantity.missing("%", "새 자료에 없음"),
        ),
    }
    texts = {
        "vehicle": PayloadText(
            key="vehicle", label_ko="차종", text=f"{NEW_VEHICLE} (synthetic)", origin="user"
        ),
        "base_date": PayloadText(key="base_date", label_ko="기준일", text=NEW_DATE, origin="user"),
        "conclusion": PayloadText(
            key="conclusion",
            label_ko="결론",
            text="B안이 중량·원가 모두 유리 (합성 자료 기준, 운영 판단 아님)",
            origin="user",
        ),
        "evidence_note": PayloadText(
            key="evidence_note", label_ko="근거 메모", text="evidence.json 참조", origin="user"
        ),
    }
    table = PayloadTable(
        name="cost_lines",
        label_ko="원가 항목(B안)",
        columns=[PayloadColumn(key="amount", label_ko="금액", unit="원")],
        rows=[
            PayloadRow(label_ko="재료비", cells={"amount": items["cost_line_material"].quantity}),
            PayloadRow(label_ko="가공비", cells={"amount": items["cost_line_process"].quantity}),
            PayloadRow(label_ko="금형비", cells={"amount": items["cost_line_tooling"].quantity}),
        ],
    )
    checks = [
        PayloadCheck(
            check_id="cost_total_sum",
            kind="sum",
            target="cost_total",
            terms=["cost_line_material", "cost_line_process", "cost_line_tooling"],
            description_ko="원가 합계 = 항목 합",
        ),
        PayloadCheck(
            check_id="cost_total_table_sum",
            kind="sum",
            target="cost_total",
            table="cost_lines",
            column="amount",
            description_ko="원가 합계 = 표 합",
        ),
        PayloadCheck(
            check_id="mass_delta_diff", kind="difference", target="mass_delta", terms=["mass_b", "mass_a"]
        ),
        PayloadCheck(
            check_id="cost_delta_diff", kind="difference", target="cost_delta", terms=["cost_b", "cost_a"]
        ),
        PayloadCheck(
            check_id="insertion_force_delta_diff",
            kind="difference",
            target="insertion_force_delta",
            terms=["insertion_force_b", "insertion_force_a"],
        ),
    ]
    return ReportPayload(
        report_id="RPT-SYNTHETIC-0001",
        created_at=created_at,
        subject=f"{NEW_VEHICLE} 설계안 A/B 비교 검토 (synthetic)",
        items=items,
        tables={"cost_lines": table},
        texts=texts,
        checks=checks,
        stale_values=[OLD_VEHICLE, OLD_DATE, OLD_COST_TEXT],
        synthetic=True,
        data_origin="synthetic",
        notes=["모든 수치는 합성 fixture 에서 계산/예측한 값이며 실제 차량·원가를 대표하지 않습니다."],
    )


def make_synthetic_documents(out_dir: str | os.PathLike[str]) -> dict[str, Path]:
    """합성 문서 fixture 일체를 out_dir 에 생성한다."""
    _require_libs()
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)
    files = {
        "past_review_2023.pptx": make_past_review_pptx(d / "past_review_2023.pptx"),
        "past_cost_table.xlsx": make_past_cost_xlsx(d / "past_cost_table.xlsx"),
        "template_review.pptx": make_template_review_pptx(d / "template_review.pptx"),
        "template_comparison.xlsx": make_template_comparison_xlsx(d / "template_comparison.xlsx"),
        "report_payload.example.json": write_payload(example_payload(), d / "report_payload.example.json"),
    }
    legacy = d / "legacy_note.ppt"
    atomic_write_bytes(
        legacy, b"synthetic placeholder - not a real binary PPT; requires approved conversion\n"
    )
    files["legacy_note.ppt"] = legacy
    atomic_write_text(d / "README.md", README_TEXT)
    files["README.md"] = d / "README.md"
    return files


if __name__ == "__main__":  # pragma: no cover - 수동 fixture 재생성
    import sys

    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("fixtures/documents")
    for name, p in make_synthetic_documents(target).items():
        print(f"{name}: {p}")
